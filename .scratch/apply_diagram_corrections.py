#!/usr/bin/env python3
"""
apply_diagram_corrections.py

1. Backs up todo/MASTER_ENI.docx.
2. Updates sections 5.7, 5.8, 5.9, 5.10, 5.11, 5.12 text to match the real architecture.
3. Replaces the 11 diagram images in the docx zip archive with the new realistic PlantUML PNGs.
4. Verifies document integrity.
"""

import os
import shutil
import zipfile
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC_DOCX = "todo/MASTER_ENI.docx"
BK_DOCX = "todo/MASTER_ENI.pre_diagram_update.docx"
BK_MEMOIRE = ".scratch/memoire_backup/MASTER_ENI.pre_diagram_update.docx"

PNG_DIR = ".scratch/diagrams_png"

# Map of images in word/media to new realistic PNGs
IMAGE_REPLACEMENTS = {
    "word/media/image70.png": os.path.join(PNG_DIR, "fig5_1_usecase.png"),
    "word/media/image68.png": os.path.join(PNG_DIR, "fig5_2_domain.png"),
    "word/media/image74.png": os.path.join(PNG_DIR, "fig5_3_sprint1.png"),
    "word/media/image76.png": os.path.join(PNG_DIR, "fig5_4_sprint3.png"),
    "word/media/image72.png": os.path.join(PNG_DIR, "fig5_5_sprint4.png"),
    "word/media/image75.png": os.path.join(PNG_DIR, "fig5_6_sprint5.png"),
    "word/media/image71.png": os.path.join(PNG_DIR, "fig5_7_sprint7.png"),
    "word/media/image77.png": os.path.join(PNG_DIR, "fig5_8_sprint9.png"),
    "word/media/image80.png": os.path.join(PNG_DIR, "fig5_9_global_classes.png"),
    "word/media/image79.png": os.path.join(PNG_DIR, "fig5_10_package.png"),
    "word/media/image78.png": os.path.join(PNG_DIR, "fig5_11_deployment.png"),
}

def set_font(run, name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

def update_text_content(doc):
    """
    Update paragraphs in sections 5.7 to 5.12 to align with real project architecture.
    """
    p_map = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}

    # Section 5.7 Intro
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Le diagramme de cas d'utilisation global présente une vue d'ensemble"):
            p.text = (
                "Le diagramme de cas d'utilisation global présente une vue d'ensemble des interactions "
                "entre les acteurs du système et les fonctionnalités offertes par l'assistant IA MIRO. "
                "Conçu pour opérer directement au sein de l'environnement Odoo 17, le système articule "
                "deux canaux majeurs d'assistance : la recherche sémantique dans la documentation d'entreprise "
                "et l'interrogation en temps réel des données transactionnelles de l'ERP. Quatre acteurs sont identifiés :"
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("l'utilisateur interne, employé de MiroKE"):
            p.text = (
                "• l'utilisateur interne, collaborateur de MiroKE (support technique, consultant ou développeur), "
                "qui sollicite l'assistant depuis l'interface Odoo pour retrouver une procédure interne ou analyser un blocage métier ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("l'utilisateur client, employé d'un client de MiroKE"):
            p.text = (
                "• l'utilisateur client, employé d'une entreprise utilisatrice accédant à l'assistant via le composant systray "
                "intégré dans sa propre instance Odoo pour consulter ses documents de gestion et l'état de ses données ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("l'administrateur système, responsable de la configuration"):
            p.text = (
                "• l'administrateur système / Odoo, responsable de la configuration des paramètres de connexion (miro.ai.settings), "
                "de la gestion du cycle de vie de l'index documentaire et de la surveillance de l'activité via le journal d'audit (miro.ai.log) ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("le système Odoo, acteur secondaire qui fournit"):
            p.text = (
                "• le système Odoo 17, acteur secondaire et tiers de confiance qui fournit le contexte d'authentification "
                "de session (uid, société active, groupes d'appartenance) et applique de manière souveraine les règles de sécurité "
                "d'enregistrement (ir.rule) et de modèle (ir.model.access) lors de chaque lecture ORM."
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Le cas d'utilisation central, Poser une question, inclut systématiquement Consulter la source"):
            p.text = (
                "Le cas d'utilisation central, Poser une question via le widget Systray OWL, inclut obligatoirement "
                "S'authentifier sur Odoo 17, garantissant que chaque échange est adossé à une identité vérifiée (uid). "
                "L'utilisateur peut facultativement Choisir le mode de source (« extend » entre « Documentation » et « Système »), "
                "laissant sinon au classifieur du middleware le soin d'orienter automatiquement la requête. Selon la nature de la question, "
                "le traitement inclut systématiquement Consulter les sources documentaires (RAG hybride BGE-M3 et BM25) ou "
                "Interroger les données ERP en temps réel (outils ORM typés en lecture seule). Enfin, le cas central inclut "
                "Visualiser le plan d'investigation et les badges de confiance, matérialisant les preuves factuelles recueillies (RG-02)."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Du côté de l'administration, le cas Indexer un document inclut"):
            p.text = (
                "Du côté de l'administration, les cas Configurer les paramètres MIRO AI et Consulter le journal d'audit et requêtes "
                "permettent de piloter le secret de signature HMAC, les timeouts et la traçabilité intégrale des échanges (RG-07). "
                "Le cas Gérer l'ingestion documentaire permet l'indexation continue des corpus PDF et wikis techniques. "
                "Le système Odoo garantit quant à lui l'étanchéité multi-sociétés en imposant la vérification des règles ir.rule lors de toute lecture de données (RG-01, RG-04)."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)

        # Section 5.8 Modèle du Domaine
        elif t.startswith("Neuf concepts métier sont identifiés :"):
            p.text = (
                "Neuf concepts métier fondamentaux structurent le domaine d'application :"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Client représente une entreprise utilisatrice de MiroKE"):
            p.text = (
                "• Société (res.company) représente l'entité juridique et organisationnelle au sein d'Odoo, délimitant "
                "strictement le périmètre des données visibles et le cloisonnement multi-sociétés ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Utilisateur représente toute personne autorisée"):
            p.text = (
                "• Utilisateur (res.users) représente l'acteur authentifié auprès d'Odoo (employé ou client), rattaché "
                "à une ou plusieurs sociétés et porteur de droits d'accès RBAC précis ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Conversation regroupe les échanges successifs"):
            p.text = (
                "• SessionConversation regroupe les échanges successifs d'un utilisateur sous un identifiant unique (session_id), "
                "conservant le mode de source sélectionné et le fil contextuel de la discussion ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Message correspond à un élément unitaire"):
            p.text = (
                "• Message correspond à un tour de parole unitaire (question utilisateur ou réponse générée par l'IA), "
                "portant son horodatage et son contenu textuel formaté en Markdown ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Source matérialise le lien entre une réponse"):
            p.text = (
                "• SourceDocumentaire matérialise la citation vérifiable d'un passage textuel précis (titre, extrait, type "
                "de document PDF/wiki, page et score de similarité) appuyant une réponse documentaire (RG-02) ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("DocumentSource représente un document brut"):
            p.text = (
                "• ObservationORM représente le résultat factuel d'une requête de lecture exécutée sur la base PostgreSQL "
                "d'Odoo (modèle cible, méthode typée, enregistrements filtrés sous l'uid demandeur) ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("IndexDocumentaire représente l'ensemble structuré"):
            p.text = (
                "• PreuveVerification (Evidence) qualifie formellement le niveau de certitude factuelle (statut, "
                "score de confiance, type de source documentaire ou système) pour interdire toute hallucination générative ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("JournalAcces trace chaque action effectuée"):
            p.text = (
                "• JournalAudit (miro.ai.log) consigne de façon immuable chaque interaction (question, réponse, durée d'exécution en ms, "
                "sources JSON, métadonnées de routage et statut), assurant la conformité et l'auditabilité (RG-07) ;"
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("DroitAccesOdoo représente la projection, au sein du domaine"):
            p.text = (
                "• ParametresMiroAI (miro.ai.settings) formalise la configuration singleton du module : URL du middleware FastAPI, "
                "secret cryptographique partagé (token HMAC), timeout de requête et activation granulaire des compétences métier."
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Le modèle met en évidence deux ramifications distinctes issues du concept central Client"):
            p.text = (
                "Le modèle met en évidence la complémentarité entre la sphère applicative Odoo (Société, Utilisateur, ParametresMiroAI, JournalAudit) "
                "et la sphère d'orchestration IA (SessionConversation, Message). Chaque Message généré par l'assistant s'adosse "
                "soit à des SourceDocumentaire issues du RAG, soit à des ObservationORM issues de la lecture sécurisée de l'ERP, "
                "l'ensemble étant systématiquement validé par une PreuveVerification avant restitution à l'utilisateur."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)

        # Section 5.9 Sprint diagrams
        elif t.startswith("Ce sprint pose les fondations de l'environnement d'inférence local. La classe GestionnaireInference"):
            p.text = (
                "Ce sprint pose les fondations de l'environnement technique d'exécution. La classe Settings centralise l'ensemble "
                "des constantes portables du système (modèle d'embedding BGE-M3, URL de l'instance locale Ollama, chemin de la base ChromaDB, "
                "tailles et recouvrements de blocs). La classe OllamaEmbeddingFunction encapsule les appels HTTP asynchrones vers le serveur "
                "d'inférence local pour générer les vecteurs denses de 1024 dimensions. Le composant FastAPIApp initialise le service web, "
                "ses gestionnaires de cycle de vie et les points d'accès de contrôle de santé."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Ce sprint introduit le cœur du pipeline d'indexation. Un DocumentSource est découpé"):
            p.text = (
                "Ce sprint implémente le pipeline d'ingestion documentaire multi-sources. La classe IngestionOrchestrator orchestre "
                "le traitement des répertoires documentaires en s'appuyant sur des chargeurs spécialisés héritant de l'interface abstraite DocumentLoader "
                "(PDFLoader pour les manuels et fiches techniques, WikiLoader pour les documentations Markdown). Les segments de texte (chunks) "
                "sont indexés au sein de ChromaManager, qui interagit avec OllamaEmbeddingFunction pour calculer et persister les vecteurs d'embeddings "
                "associés à leurs métadonnées d'origine."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Ce sprint introduit le moteur de recherche proprement dit. Une RequeteRecherche est traitée"):
            p.text = (
                "Ce sprint réalise le moteur de recherche documentaire hybride. Une RequeteRecherche est prise en charge par la classe "
                "SearchHybrid, qui exécute en parallèle une recherche sémantique dense sur ChromaManager et une recherche lexicale inversée "
                "sur BM25Index. Les deux flux de résultats sont fusionnés au moyen de l'algorithme Reciprocal Rank Fusion (RRF), pondéré par les "
                "coefficients de pertinence RAG_VECTOR_WEIGHT (0.6) et RAG_BM25_WEIGHT (0.4), afin de produire les passages documentaires les plus fiables (ResultatFiltre)."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Ce sprint clôt la chaîne de traitement du pipeline RAG. L'AssembleurPrompt"):
            p.text = (
                "Ce sprint clôt la chaîne de traitement RAG et introduit les gardes anti-hallucination. La classe GenerateNode assemble "
                "le gabarit de prompt enrichi des passages contextuels sourcés et sollicite le LLM local. La réponse brute est ensuite transmise "
                "à la classe VerifyAnswerNode, qui vérifie que chaque affirmation est rigoureusement soutenue par les citations fournies et génère "
                "un VerificationReport. L'objet final AssistantResponse garantit ainsi que chaque élément de réponse est accompagné de ses SourceReference vérifiées."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Ce sprint introduit les classes nécessaires à l'accès à l'assistant depuis l'instance Odoo"):
            p.text = (
                "Ce sprint introduit l'intégration native au sein du client web Odoo 17. Le composant frontend MiroAiSystray, développé en OWL 2.0, "
                "fournit l'interface de conversation accessible depuis la barre supérieure, gère l'état réactif et le basculement du mode de source. "
                "Il dialogue par requêtes RPC avec MiroAiController. Ce dernier délègue à AskFlow la génération d'un jeton JWT signé par HMAC SHA-256 "
                "et l'émission d'une clé API éphémère à usage unique, relayée ensuite au middleware FastAPI par BackendClient selon un schéma de scrutation asynchrone (Job Polling)."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Ce sprint introduit les classes garantissant le respect des règles de gestion RG-01 et RG-04. Le Ges"):
            p.text = (
                "Ce sprint concrétise la sécurité multi-sociétés, la lecture des données ERP et l'auditabilité. Le modèle Odoo MiroAiSettings "
                "gère les paramètres et secrets d'authentification, tandis que MiroAiLog consigne chaque requête de manière immuable. Pour l'accès aux données, "
                "la classe OrmRuntimeReadTools expose des méthodes d'inspection et de lecture strictement typées (orm_search_read, orm_search_count, orm_read_group), "
                "exécutées à travers OdooConnector sous l'identité exacte de l'utilisateur demandeur (uid) sans jamais contourner les règles ir.rule ni recourir à sudo()."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)

        # Section 5.10 Global Class Diagram
        elif t.startswith("Cette couche regroupe les classes responsables de l'accès au système depuis l'instance Odoo d'un cli"):
            p.text = (
                "Cette couche regroupe les composants Odoo 17 : le widget réactif MiroAiSystray (OWL 2.0), le contrôleur HTTP MiroAiController, "
                "le module de signature AskFlow et les modèles de persistance Odoo MiroAiSettings et MiroAiLog."
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Cette couche porte la responsabilité du cloisonnement entre clients et du contrôle d'accès. Le Clien"):
            p.text = (
                "Cette couche middleware héberge le point d'entrée FastAPIApp, le gestionnaire de tâches asynchrones AskJobsManager "
                "et le graphe d'orchestration LangGraphPipeline, qui oriente la requête via RouteNode et synthétise l'objet AssistantResponse."
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Cette couche constitue le cœur fonctionnel du pipeline RAG. Une Conversation regroupe des Message, c"):
            p.text = (
                "Cette couche encapsule la recherche et l'indexation documentaire : IngestionOrchestrator, ChromaManager, BM25Index, "
                "le moteur SearchHybrid et les nœuds décisionnels GenerateNode et VerifyAnswerNode."
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Cette couche assure la transformation des documents bruts en représentations exploitables par le mot"):
            p.text = (
                "Cette couche assure l'interrogation en temps réel des données de gestion d'Odoo : l'agent ReAct OrmAgentLoop, "
                "la boîte à outils de lecture seule OrmRuntimeReadTools et le connecteur XML-RPC sécurisé OdooConnector."
            )
            for r in p.runs:
                set_font(r, size_pt=12)
        elif t.startswith("Cette couche, la plus stable du système, gère le cycle de vie du modèle de langage local. Le Gestion"):
            p.text = (
                "Les dépendances entre couches respectent strictement le principe de moindre privilège et de découplage : "
                "le contrôleur Odoo sollicite le middleware FastAPI par requêtes signées asynchrones (HTTP 202), le middleware active la branche documentaire "
                "ou la branche système, et les outils ORM interrogent l'instance Odoo sous l'identité vérifiée du demandeur."
            )
            for r in p.runs:
                set_font(r, size_pt=12)

        # Section 5.11 Package Diagram
        elif t.startswith("Le paquetage Interface & Intégration dépend du paquetage Sécurité & Multi-tenant pour vérifier les d"):
            p.text = (
                "Le paquetage odoo_addons.miro_ai dialogue avec le paquetage api via des requêtes HTTP REST asynchrones signées par HMAC SHA-256. "
                "Le paquetage api confie le traitement au graphe d'orchestration de core, qui sollicite à son tour document_rag pour la recherche "
                "documentaire ou erp_data pour la consultation des données du progiciel. Les paquetages core, document_rag et erp_data s'appuient "
                "sur config pour charger les paramètres et adresses d'accès sans duplication de configuration. Enfin, erp_data réinterroge l'instance Odoo "
                "via XML-RPC sous l'identité authentifiée de l'utilisateur."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)

        # Section 5.12 Deployment Diagram
        elif t.startswith("Ce composant central, développé avec FastAPI, orchestre les échanges entre les trois autres services"):
            p.text = (
                "Le serveur IA dédié souverain héberge sous conteneurs Docker le middleware FastAPI / LangGraph (port 8000), le serveur d'inférence "
                "local Ollama / vLLM (port 11434) exploitant la carte GPU NVIDIA RTX, ainsi que la base vectorielle ChromaDB persistée sur disque local. "
                "L'instance Odoo 17 et sa base relationnelle PostgreSQL 16 (port 5432) résident sur le serveur ERP de l'entreprise. "
                "Le poste utilisateur exécute quant à lui le composant OWL 2.0 dans son navigateur web, communiquant avec Odoo via HTTPS (port 8069). "
                "Toutes les communications entre Odoo et le serveur IA transitent par le réseau local interne sécurisé, sans aucune dépendance envers un cloud externe."
            )
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_font(r, size_pt=12)

def main():
    print(f"Creating backup of {SRC_DOCX}...")
    shutil.copyfile(SRC_DOCX, BK_DOCX)
    os.makedirs(os.path.dirname(BK_MEMOIRE), exist_ok=True)
    shutil.copyfile(SRC_DOCX, BK_MEMOIRE)
    print(f"Backups created at {BK_DOCX} and {BK_MEMOIRE}")

    print("Step 1: Updating text in sections 5.7 to 5.12...")
    doc = docx.Document(SRC_DOCX)
    update_text_content(doc)
    doc.save(SRC_DOCX)
    print("Text content updated successfully.")

    print("Step 2: Replacing diagram images in docx zip archive...")
    temp_zip = "/tmp/temp_replaced_doc.zip"
    with zipfile.ZipFile(SRC_DOCX, "r") as zin, zipfile.ZipFile(temp_zip, "w") as zout:
        for item in zin.infolist():
            if item.filename in IMAGE_REPLACEMENTS:
                rep_path = IMAGE_REPLACEMENTS[item.filename]
                with open(rep_path, "rb") as f:
                    zout.writestr(item, f.read())
                print(f"  Replaced {item.filename} with {rep_path}")
            else:
                zout.writestr(item, zin.read(item.filename))

    os.replace(temp_zip, SRC_DOCX)
    print("All 11 diagram images successfully replaced in docx archive.")

    print("Step 3: Verification of the final document...")
    doc_final = docx.Document(SRC_DOCX)
    print(f"Final paragraphs: {len(doc_final.paragraphs)}")
    print(f"Final tables: {len(doc_final.tables)}")

    # Check Chapters 1 & 2
    doc_bk = docx.Document(BK_DOCX)
    diffs = []
    for i in range(226, 355):
        t1 = doc_final.paragraphs[i].text
        t2 = doc_bk.paragraphs[i].text
        if t1 != t2:
            diffs.append((i, t1[:50], t2[:50]))
    print(f"Differences in Chapters 1 & 2 (p. 226-354): {len(diffs)}")
    assert len(diffs) == 0, "Chapters 1 & 2 were modified!"

    print("\nSUCCESS: Document updated and verified perfectly!")

if __name__ == "__main__":
    main()
