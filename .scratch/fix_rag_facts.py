import copy
from docx import Document

doc = Document('todo/MASTER_ENI.docx')


def set_cell(cell, text):
    p = cell.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find_row(label, col=0):
    hits = [row for t in doc.tables for row in t.rows if row.cells[col].text.strip().startswith(label)]
    return hits


rrf = [r for r in find_row('Pondération RRF')]
assert len(rrf) == 1
set_cell(rrf[0].cells[1], 'k = 60 ; poids 0,7 (dense) et 0,3 (BM25)')
fusion = [r for r in find_row('Fusion') if 'Reciprocal' in r.cells[1].text]
assert len(fusion) == 1
set_cell(fusion[0].cells[2], 'Poids 0,7 (vectoriel) et 0,3 (lexical)')
returned = find_row('Passages retournés')
assert len(returned) == 1
set_cell(returned[0].cells[1], '10 dans le graphe principal')
set_cell(returned[0].cells[2], 'Candidats : 20 fois ce nombre, au moins 100')

p751 = next(p for p in doc.paragraphs if p.text.startswith('Le modèle de génération reçoit les passages numérotés'))
run = p751.runs[0]
old = ("Une vérification déterministe contrôle ensuite que les éléments des listes et tableaux des passages cités figurent bien dans la réponse, et calcule un taux de restitution.")
assert old in run.text
run.text = run.text.replace(old, "Un module de contrôle de la complétude de la restitution existe dans le dépôt, mais il n'est pas branché sur le chemin de production : son intégration reste à faire.")

body = copy.deepcopy(p751._p)
placeholder = next(p for p in doc.paragraphs if p.text.startswith('[À compléter : résultat de tests/eval_rag.py'))
placeholder._p.addprevious(body)
new = next(p for p in doc.paragraphs if p._p is body)
new.runs[0].text = ("Le banc de recherche seule (scripts/benchmark_rag.py) donne 10/10 de réussites et un Hit@1 de 9/9 depuis la correction, "
                    "le 27 octobre 2026, du découpage des mots accentués dans l'index lexical, qui coupait par exemple « déploiement » "
                    "en « ploiement ».")
for r in new.runs[1:]:
    r._r.getparent().remove(r._r)

bullets = [p for p in doc.paragraphs if p.text.startswith("• chaque question a été rejouée une seule fois")]
assert len(bullets) == 1
clone = copy.deepcopy(bullets[0]._p)
bullets[0]._p.addnext(clone)
extra = next(p for p in doc.paragraphs if p._p is clone)
extra.runs[-1].text = ("la recherche documentaire n'applique pas de seuil de pertinence absolu : la fusion normalise les scores par "
                       "le maximum, si bien que des passages hors sujet peuvent atteindre le modèle (constat du 27 octobre 2026)")
doc.save('todo/MASTER_ENI.docx')
print('ok')
