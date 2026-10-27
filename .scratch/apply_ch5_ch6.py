#!/usr/bin/env python3
"""
Script to apply Chapter 5 repositioning and Chapter 6 insertion on todo/MASTER_ENI.docx.
Creates backups before proceeding.
"""

import os
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_run_font(run, name='Times New Roman', size_pt=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def create_styled_paragraph(doc, text="", style='Normal', bold=False, italic=False, size_pt=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.style = style
    p.paragraph_format.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_run_font(r, size_pt=size_pt, bold=bold, italic=italic)
    return p

def create_subheading(doc, text, level=3, size_pt=14):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=size_pt, bold=True)
    return p

def create_bullet(doc, text, prefix="", size_pt=12):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_bullet = p.add_run("• ")
    set_run_font(r_bullet, size_pt=size_pt, bold=True)
    
    if prefix:
        r_pre = p.add_run(prefix)
        set_run_font(r_pre, size_pt=size_pt, bold=True)
        
    r_text = p.add_run(text)
    set_run_font(r_text, size_pt=size_pt)
    return p

def create_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.style = 'Table10'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for idx, title in enumerate(headers):
        hdr_cells[idx].text = title
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            set_run_font(r, size_pt=10.5, bold=True)
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0 and len(str(cell_value)) <= 8:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 1 and len(str(cell_value)) <= 6:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                set_run_font(r, size_pt=10)
                
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    return tbl

def execute_reposition_and_ch6(doc_path):
    print(f"Loading document: {doc_path}...")
    doc = docx.Document(doc_path)
    body = doc._body._element

    old_ch5_h = None
    new_ch5_h = None
    conclusion_h = None

    for idx, child in enumerate(body):
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            t = p.text.strip().upper()
            if 'CHAPITRE 5' in t and p.style.name == 'Heading 2':
                if old_ch5_h is None:
                    old_ch5_h = (idx, child, p.text)
                elif new_ch5_h is None:
                    new_ch5_h = (idx, child, p.text)
            elif 'CONCLUSION' in t and p.style.name == 'Heading 1' and conclusion_h is None:
                conclusion_h = (idx, child, p.text)

    assert old_ch5_h[0] == 501, f"Expected old Ch5 at 501, got {old_ch5_h[0]}"
    assert new_ch5_h[0] == 1083, f"Expected new Ch5 at 1083, got {new_ch5_h[0]}"
    assert conclusion_h[0] == 1281, f"Expected conclusion at 1281, got {conclusion_h[0]}"

    print("Found exact boundaries:")
    print(f"  Old Chapter 5 at BodyElem {old_ch5_h[0]}")
    print(f"  New Chapter 5 at BodyElem {new_ch5_h[0]}")
    print(f"  Conclusion at BodyElem {conclusion_h[0]}")

    # 1. Collect all elements of user's new Chapter 5
    new_ch5_elements = []
    curr = new_ch5_h[1]
    while curr is not None and curr != conclusion_h[1]:
        new_ch5_elements.append(curr)
        curr = curr.getnext()

    # 2. Collect legacy elements to delete: between old_ch5_h and new_ch5_h
    legacy_elements = []
    curr = old_ch5_h[1]
    while curr is not None and curr != new_ch5_h[1]:
        legacy_elements.append(curr)
        curr = curr.getnext()

    print(f"Moving {len(new_ch5_elements)} elements of Chapter 5...")
    for elem in new_ch5_elements:
        old_ch5_h[1].addprevious(elem)

    print(f"Removing {len(legacy_elements)} legacy Bachelor elements...")
    for elem in legacy_elements:
        elem.getparent().remove(elem)

    # 3. Format and update captions and tables in Chapter 5
    # The tables in Chapter 5 are tables 16 to 32 (17 tables)
    ch5_captions_map = [
        (16, "Tableau 16. Synthèse des rôles de l'équipe SCRUM", "replace"),
        (17, "Tableau 17. Découpage en releases du projet", "insert"),
        (18, "Tableau 18. Product Backlog — Release 1 (Infrastructure et modèles)", "replace"),
        (19, "Tableau 19. Product Backlog — Release 2 (Pipeline RAG et données)", "replace"),
        (20, "Tableau 20. Product Backlog — Release 3 (Interface utilisateur et intégration)", "replace"),
        (21, "Tableau 21. Product Backlog — Release 4 (Sécurité, tests et documentation)", "replace"),
        (22, "Tableau 22. Synthèse des estimations d'effort du Product Backlog", "replace"),
        (23, "Tableau 23. Planification des sprints par release", "insert"),
        (24, "Tableau 24. Sprint Backlog — Sprint 3 (Indexation des documents internes)", "replace"),
        (25, "Tableau 25. Sprint Backlog — Sprint 4 (Recherche sémantique et code source)", "replace"),
        (26, "Tableau 26. Sprint Backlog — Sprint 5 (Génération de réponses sourcées)", "replace"),
        (27, "Tableau 27. Dictionnaire des données — Entité Client", "replace"),
        (28, "Tableau 28. Dictionnaire des données — Entité Conversation", "replace"),
        (29, "Tableau 29. Dictionnaire des données — Entité Document_Source", "replace"),
        (30, "Tableau 30. Dictionnaire des données — Entité Journal_Acces", "replace"),
        (31, "Tableau 31. Dictionnaire des données — Entité Message", "replace"),
        (32, "Tableau 32. Dictionnaire des données — Entité Utilisateur", "replace"),
    ]

    for tbl_idx, caption_text, mode in ch5_captions_map:
        if tbl_idx < len(doc.tables):
            t = doc.tables[tbl_idx]
            t.style = 'Table10'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Format header row
            for cell in t.rows[0].cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for r in p.runs:
                        set_run_font(r, size_pt=10, bold=True)
            # Format data rows
            for row in t.rows[1:]:
                for c_i, cell in enumerate(row.cells):
                    for p in cell.paragraphs:
                        if c_i == 0 and len(cell.text.strip()) <= 8:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        for r in p.runs:
                            set_run_font(r, size_pt=9.5)

            # Update or insert caption paragraph
            prev_elem = t._tbl.getprevious()
            if mode == "replace" and prev_elem is not None and prev_elem.tag.split('}')[-1] == 'p':
                prev_p = docx.text.paragraph.Paragraph(prev_elem, doc)
                prev_p.text = caption_text
                prev_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                prev_p.paragraph_format.space_before = Pt(8)
                prev_p.paragraph_format.space_after = Pt(4)
                prev_p.paragraph_format.keep_with_next = True
                for r in prev_p.runs:
                    set_run_font(r, size_pt=11, bold=True)
            elif mode == "insert":
                p_new = doc.add_paragraph()
                p_new.style = 'Normal'
                p_new.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_new.paragraph_format.space_before = Pt(8)
                p_new.paragraph_format.space_after = Pt(4)
                p_new.paragraph_format.keep_with_next = True
                r = p_new.add_run(caption_text)
                set_run_font(r, size_pt=11, bold=True)
                t._tbl.addprevious(p_new._p)

    print("Chapter 5 tables successfully styled with Table10 and numbered captions.")

    # 4. Insert Chapitre 6 (Intégration IA dans la conception) right before conclusion_h[1]
    ref = conclusion_h[1]

    # Heading Chapitre 6
    p_h6 = doc.add_paragraph()
    p_h6.style = 'Heading 2'
    p_h6.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_h6.paragraph_format.space_before = Pt(18)
    p_h6.paragraph_format.space_after = Pt(8)
    r_h6 = p_h6.add_run("Chapitre 6. Intégration IA dans la conception")
    set_run_font(r_h6, size_pt=16, bold=True)
    ref.addprevious(p_h6._p)

    # Intro Chapitre 6
    p_ch6_intro = create_styled_paragraph(doc, "L'introduction de capacités d'intelligence artificielle au sein d'un Progiciel de Gestion Intégré tel qu'Odoo 17 ne se limite pas à l'appel distant d'un modèle de langage. Elle exige une conception d'ingénierie rigoureuse afin d'articuler harmonieusement les flux de données hétérogènes, de garantir la souveraineté des traitements et de maintenir le contrôle d'accès strict propre au système d'information de l'entreprise. Ce chapitre formalise l'intégration de l'IA dans le système : la définition détaillée des modules composants l'assistant MIRO AI, la caractérisation des flux de données nécessaires, la justification technique des modèles retenus (embedding dense et LLM de génération) et enfin les modalités d'intégration par microservice asynchrone.", space_before=6)
    ref.addprevious(p_ch6_intro._p)

    # 6.1. Définition du module IA
    h6_1 = create_subheading(doc, "6.1. Définition du module IA", size_pt=14)
    ref.addprevious(h6_1._p)

    p6_1_1 = create_styled_paragraph(doc, "Le système MIRO AI est conçu selon une architecture modulaire et découplée articulée autour de deux composants logiciels fondamentaux en interaction continue :")
    ref.addprevious(p6_1_1._p)

    mod_list = [
        ("Le module client Odoo 17 (« miro_ai ») : ", "Développé nativement en JavaScript OWL 2.0 (Odoo Web Library) et Python Odoo, ce composant s'insère directement dans le client web d'Odoo. Il fournit une interface conversationnelle ergonomique accessible depuis la barre supérieure (systray), supporte le rendu dynamique du Markdown, gère l'historique des échanges et affiche les badges de confiance et les sources citées. Côté serveur Odoo, un contrôleur asynchrone sécurisé relaie les requêtes des utilisateurs vers le moteur IA en y associant l'identité stricte (uid), la société active (company_id) et les groupes de privilèges, tout en assurant la traçabilité complète dans le journal d'audit miro.ai.log."),
        ("Le serveur middleware d'orchestration IA (FastAPI & LangGraph) : ", "Service backend autonome déployé sur serveur sécurisé privé. Il héberge un graphe décisionnel LangGraph chargé de classer la question, d'orienter la requête vers le canal approprié (recherche documentaire ou consultation système ERP) et de piloter les moteurs d'exécution spécialisés."),
    ]
    for pre, txt in mod_list:
        p_mod = create_bullet(doc, txt, prefix=pre)
        ref.addprevious(p_mod._p)

    p6_1_2 = create_styled_paragraph(doc, "Au sein du backend d'orchestration, deux moteurs spécialisés opèrent selon la nature de la sollicitation :")
    ref.addprevious(p6_1_2._p)

    moteurs_list = [
        ("Le moteur RAG documentaire hybride : ", "Dédié à la recherche dans les connaissances d'entreprise non structurées (manuels de procédures, guides techniques, wikis internes, politiques de gestion). Il combine une recherche sémantique dense (embeddings BGE-M3) et une recherche lexicale exacte (Rank-BM25), fusionnées par l'algorithme Reciprocal Rank Fusion (RRF). Un module d'assemblage et de vérification anti-hallucination contraint le LLM à étayer chaque affirmation par une citation explicite des passages documentaires récupérés."),
        ("L'agent d'accès et de lecture ORM en temps réel : ", "Spécifiquement conçu pour l'interrogation des données transactionnelles et de gestion d'Odoo. Cet agent exploite un ensemble d'outils typés en lecture seule (inspection du schéma, comptage, listage et agrégation d'indicateurs). Il exécute impérativement chaque requête sous le contexte de sécurité authentifié du demandeur (uid), respectant scrupuleusement les droits de lecture (ir.model.access) et les filtres de cloisonnement multi-sociétés (ir.rule) sans jamais recourir à sudo() ni à des requêtes SQL brutes non contrôlées."),
    ]
    for pre, txt in moteurs_list:
        p_mot = create_bullet(doc, txt, prefix=pre)
        ref.addprevious(p_mot._p)

    # 6.2. Données nécessaires
    h6_2 = create_subheading(doc, "6.2. Données nécessaires au système", size_pt=14)
    ref.addprevious(h6_2._p)

    p6_2_1 = create_styled_paragraph(doc, "Le bon fonctionnement de MIRO AI repose sur l'exploitation coordonnée de trois catégories de données distinctes mais complémentaires :")
    ref.addprevious(p6_2_1._p)

    data_types = [
        ("Données documentaires non structurées (Corpus de connaissances) : ", "Fichiers PDF, pages de wikis, documentations techniques Markdown, comptes rendus opérationnels et tickets de support archivés. Ces documents font l'objet d'un pré-traitement rigoureux : extraction textuelle, nettoyage des caractères parasites, segmentation sémantique en blocs homogènes (chunks de 512 tokens avec un recouvrement de 64 tokens pour préserver la continuité du contexte) et vectorisation."),
        ("Données relationnelles structurées et métadonnées de l'ERP Odoo : ", "Schéma relationnel de la base PostgreSQL hébergeant les tables de gestion d'Odoo 17 (res_partner, sale_order, account_move, stock_picking, hr_leave) ainsi que les tables système de sécurité et de structure (ir_model, ir_model_fields, ir_model_access, ir_rule). Ces données ne sont jamais dupliquées dans l'IA ; elles sont consultées à la volée via l'ORM sous le profil de l'utilisateur connecté."),
        ("Données vectorielles et d'indexation sémantique : ", "Vecteurs numériques denses de 1024 dimensions générés par le modèle BGE-M3 pour chaque bloc documentaire, indexés au sein d'une base vectorielle confinée en local (ChromaDB / Qdrant), parallèlement à un index inversé de fréquences de termes pour l'algorithme BM25."),
    ]
    for pre, txt in data_types:
        p_dt = create_bullet(doc, txt, prefix=pre)
        ref.addprevious(p_dt._p)

    p6_2_2 = create_styled_paragraph(doc, "Le Tableau 33 synthétise la matrice des flux informationnels entre les différents magasins de données et les modules d'IA.")
    ref.addprevious(p6_2_2._p)

    # Caption Table 33
    p_cap33 = create_styled_paragraph(doc, "Tableau 33. Matrice des flux et sources de données exploitées par le système", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap33._p)

    headers_tbl33 = ["Source de données", "Type de contenu", "Mode d'accès", "Rôle dans le système"]
    data_tbl33 = [
        ["Documentation interne (PDF, MD)", "Texte non structuré", "Indexation vectorielle et lexicale (RAG)", "Répondre aux questions procédurales et fonctionnelles"],
        ["Base relationnelle Odoo (PostgreSQL)", "Données métier transactionnelles", "ORM Odoo temps réel sous uid strict", "Fournir l'état actuel des ventes, stocks et factures"],
        ["Métadonnées Odoo (ir.model, ir.rule)", "Définition du schéma et règles RBAC", "Inspection de métadonnées sans sudo()", "Valider la visibilité et guider la construction des requêtes"],
        ["Base vectorielle locale (ChromaDB)", "Embeddings 1024d + métadonnées", "Recherche par similarité cosinus (top-k)", "Sélection des passages les plus proches de la requête"],
        ["Journal d'audit (miro.ai.log)", "Logs des requêtes, durées et sources", "Écriture asynchrone par le contrôleur", "Assurer la conformité, la traçabilité et l'auditabilité"],
    ]
    tbl33 = create_styled_table(doc, headers_tbl33, data_tbl33, col_widths=[1.8, 1.8, 1.8, 2.1])
    ref.addprevious(tbl33._tbl)

    p_sp7 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp7._p)

    # 6.3. Choix du modèle IA et justification
    h6_3 = create_subheading(doc, "6.3. Choix du modèle IA et justification", size_pt=14)
    ref.addprevious(h6_3._p)

    p6_3_intro = create_styled_paragraph(doc, "Dans une architecture d'IA souveraine opérant en circuit fermé sur les serveurs de Miro K.E., le choix des modèles d'embedding et des modèles de fondation génératifs obéit à un compromis strict entre qualité du raisonnement, fidélité aux sources, latence d'inférence et contrainte de mémoire VRAM.")
    ref.addprevious(p6_3_intro._p)

    h_emb = create_subheading(doc, "Modèle d'embedding sémantique retenu : BGE-M3 (BAAI)", size_pt=12)
    ref.addprevious(h_emb._p)

    p_emb = create_styled_paragraph(doc, "Le modèle BGE-M3 (BAAI/bge-m3) a été sélectionné pour la représentation vectorielle des connaissances d'entreprise. Il présente des avantages déterminants pour notre cas d'usage :")
    ref.addprevious(p_emb._p)

    emb_points = [
        ("Polyvalence multi-représentation : ", "BGE-M3 est capable de générer simultanément des embeddings denses (captant la sémantique globale), des poids lexicaux clairsemés (assimilables au BM25 pour les mots-clés rares) et des représentations multi-vecteurs (ColBERT) pour un alignement précis token à token."),
        ("Support multilingue de premier plan : ", "Entraîné sur plus de 100 langues, il offre une excellente compréhension du français professionnel, des terminologies juridico-comptables et des acronymes techniques propres à Odoo."),
        ("Fenêtre de contexte étendue : ", "Capable de traiter des séquences atteignant 8 192 tokens, évitant la troncation brutale des sections de manuels techniques complexes."),
    ]
    for pre, txt in emb_points:
        p_ep = create_bullet(doc, txt, prefix=pre)
        ref.addprevious(p_ep._p)

    h_llm = create_subheading(doc, "Modèles de langage pour la génération locale (LLM)", size_pt=12)
    ref.addprevious(h_llm._p)

    p_llm = create_styled_paragraph(doc, "Pour la synthèse des réponses et le raisonnement agentique, plusieurs modèles ouverts de nouvelle génération ont été évalués sur l'infrastructure matérielle locale de Miro K.E. Le Tableau 34 présente l'analyse comparative des modèles candidats.")
    ref.addprevious(p_llm._p)

    # Caption Table 34
    p_cap34 = create_styled_paragraph(doc, "Tableau 34. Analyse comparative des modèles de langage candidats pour l'inférence locale", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap34._p)

    headers_tbl34 = ["Modèle candidat", "Paramètres", "VRAM requise (quantifié 4/8-bit)", "Fidélité au contexte (Anti-hallucination)", "Qualité du français", "Verdict et décision"]
    data_tbl34 = [
        ["Qwen 2.5 (14B)", "14 milliards", "9 à 11 Go", "Excellente (suivi strict des instructions)", "Très élevée", "Retenu : Modèle principal pour les requêtes complexes et l'ORM"],
        ["Llama 3.1 (8B)", "8 milliards", "5 à 6 Go", "Très bonne", "Bonne", "Retenu : Modèle rapide pour la classification et le RAG léger"],
        ["Mistral (7B v0.3)", "7 milliards", "5 à 6 Go", "Moyenne (tendance à la prolixité)", "Excellente", "Alternative de repli pour la recherche documentaire"],
        ["Phi-3.5 Mini (3.8B)", "3.8 milliards", "3 Go", "Moyenne sur raisonnements ERP complexes", "Moyenne", "Écarté : Capacités de raisonnement multi-tables insuffisantes"],
    ]
    tbl34 = create_styled_table(doc, headers_tbl34, data_tbl34, col_widths=[1.5, 1.0, 1.4, 1.5, 1.0, 1.6])
    ref.addprevious(tbl34._tbl)

    p_sp8 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp8._p)

    p_justif_llm = create_styled_paragraph(doc, "Le modèle Qwen 2.5 (14B) quantifié en format GGUF Q4_K_M est retenu comme moteur génératif de référence. Il combine une capacité de raisonnement et de fidélité aux instructions remarquable, un excellent traitement de la langue française et une exécution fluide sur notre serveur équipé d'une carte GPU NVIDIA RTX (tenant dans 12 Go de VRAM). Pour les environnements de test plus contraints, Llama 3.1 (8B) sert de solution alternative hautement réactive.")
    ref.addprevious(p_justif_llm._p)

    p_cap35 = create_styled_paragraph(doc, "Tableau 35. Paramétrage des hyperparamètres d'indexation et d'inférence", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap35._p)

    headers_tbl35 = ["Paramètre", "Valeur retenue", "Justification technique"]
    data_tbl35 = [
        ["Taille des chunks (Chunk size)", "512 tokens", "Granularité optimale pour isoler une règle métier sans bruit contextuel"],
        ["Recouvrement (Chunk overlap)", "64 tokens", "Préservation de la continuité syntaxique entre paragraphes adjacents"],
        ["Top-k documents récupérés", "5 passages", "Couverture suffisante pour la synthèse sans saturer la fenêtre de contexte"],
        ["Pondération RRF (Dense vs BM25)", "k = 60 (RRF standard)", "Fusion robuste équilibrant sémantique et correspondance lexicale exacte"],
        ["Température de génération (LLM)", "0.1 à 0.2", "Génération quasi déterministe pour bannir toute extrapolation créative"],
        ["Top-p (Nucleus sampling)", "0.9", "Filtrage des jetons à faible probabilité pour une précision factuelle maximale"],
    ]
    tbl35 = create_styled_table(doc, headers_tbl35, data_tbl35, col_widths=[2.2, 1.8, 3.5])
    ref.addprevious(tbl35._tbl)

    p_sp9 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp9._p)

    # 6.4. Méthode d’intégration
    h6_4 = create_subheading(doc, "6.4. Méthode d’intégration logicielle", size_pt=14)
    ref.addprevious(h6_4._p)

    p6_4_1 = create_styled_paragraph(doc, "L'intégration entre l'ERP Odoo 17 et la couche d'intelligence artificielle repose sur un modèle d'architecture en microservice découplé, interconnecté par des API REST sécurisées et asynchrones. Ce choix architectural répond à des impératifs critiques :")
    ref.addprevious(p6_4_1._p)

    integ_reasons = [
        ("Isolation des ressources de calcul : ", "L'inférence des grands modèles de langage sollicite intensivement la mémoire GPU et le processeur. Découpler le moteur IA sous un service FastAPI distinct empêche tout risque de saturation des workers HTTP standards d'Odoo, garantissant la fluidité ininterrompue des opérations de gestion quotidiennes."),
        ("Sécurité et authenticité des flux (HMAC SHA-256) : ", "Chaque requête transmise par le contrôleur Odoo au backend IA est signée cryptographiquement à l'aide d'un secret partagé (signing_secret) configuré dans les paramètres de miro_ai. Une clé de requête éphémère (requester_key) à usage unique est générée par job, interdisant toute attaque par rejeu ou accès externe non authentifié."),
        ("Pattern de traitement asynchrone (Job Polling) : ", "En raison du délai incompressible requis par l'inférence locale (de 2 à 5 secondes selon la complexité), la requête est traitée de manière non bloquante. Le contrôleur Odoo soumet la tâche et reçoit immédiatement un identifiant unique (job_id) avec le code HTTP 202 Accepted. Le composant OWL de l'interface interroge ensuite l'état du job toutes les 2 secondes jusqu'à la mise à disposition de la réponse complète."),
        ("Confinement réseau absolu : ", "Les communications s'opèrent exclusivement au sein du réseau local d'entreprise ou du réseau virtuel Docker interne, sans aucune exposition sur Internet."),
    ]
    for pre, txt in integ_reasons:
        p_ir = create_bullet(doc, txt, prefix=pre)
        ref.addprevious(p_ir._p)

    p6_4_2 = create_styled_paragraph(doc, "Le Tableau 36 détaille le contrat d'interface des points d'accès exposés entre le module Odoo miro_ai et le middleware FastAPI.")
    ref.addprevious(p6_4_2._p)

    # Caption Table 36
    p_cap36 = create_styled_paragraph(doc, "Tableau 36. Spécification des points d'accès API (Endpoints) entre miro_ai et le backend IA", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap36._p)

    headers_tbl36 = ["Méthode & Route", "Émetteur -> Récepteur", "Données transmises", "Réponse et fonction opérationnelle"]
    data_tbl36 = [
        ["POST /api/v1/health", "Odoo -> FastAPI", "Token HMAC de contrôle", "HTTP 200 {status: ok} : Vérification de la disponibilité du serveur IA"],
        ["POST /api/v1/ask/jobs", "Odoo -> FastAPI", "question, uid, company_id, groups, session_id, HMAC", "HTTP 202 {job_id, status: pending} : Lancement asynchrone de l'analyse"],
        ["GET /api/v1/ask/jobs/{id}", "Odoo -> FastAPI", "job_id, headers HMAC", "HTTP 200 {status: done, answer, sources, confidence} : Restitution du résultat"],
        ["POST /api/v1/sessions", "Odoo -> FastAPI", "uid, limit", "HTTP 200 {sessions: [...]} : Liste des conversations récentes de l'utilisateur"],
        ["POST /miro_ai/ask", "Web Client OWL -> Odoo", "question, session_id, source_mode", "Relais contrôleur Odoo : initie le job et retourne le job_id au widget"],
        ["POST /miro_ai/ask_result", "Web Client OWL -> Odoo", "job_id, session_id", "Scrutation périodique : retourne pending ou la réponse formatée avec sources"],
    ]
    tbl36 = create_styled_table(doc, headers_tbl36, data_tbl36, col_widths=[2.0, 1.5, 1.8, 2.2])
    ref.addprevious(tbl36._tbl)

    p_sp10 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp10._p)

    # Add PARTIE III heading before conclusion
    p_p3 = doc.add_paragraph()
    p_p3.style = 'Heading 1'
    p_p3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_p3.paragraph_format.space_before = Pt(24)
    p_p3.paragraph_format.space_after = Pt(12)
    r_p3 = p_p3.add_run("PARTIE III : RÉALISATION")
    set_run_font(r_p3, size_pt=18, bold=True)
    ref.addprevious(p_p3._p)

    p_p3_intro = create_styled_paragraph(doc, "Après avoir mené l'analyse préalable des besoins et formalisé la conception conceptuelle et technique des composants d'intelligence artificielle, cette troisième partie est consacrée à la concrétisation et au déploiement du système au sein de l'écosystème Odoo 17. Elle s'articule autour de trois axes : la mise en place de l'environnement de développement et de déploiement conteneurisé (Chapitre 7), le développement effectif du module Odoo miro_ai et des moteurs RAG et ORM (Chapitre 8), ainsi que les protocoles de tests, de validation fonctionnelle et d'évaluation des performances (Chapitre 9).", space_before=6)
    ref.addprevious(p_p3_intro._p)

    print("Repositioning and Chapter 6 insertion complete.")
    print(f"Saving to {doc_path}...")
    doc.save(doc_path)
    print("Saved successfully!")

if __name__ == '__main__':
    src = 'todo/MASTER_ENI.docx'
    bk1 = 'todo/MASTER_ENI.pre_ch5_ch6.docx'
    bk2 = '.scratch/memoire_backup/MASTER_ENI.pre_ch5_ch6.docx'
    
    print(f"Creating backup copies to {bk1} and {bk2}...")
    shutil.copyfile(src, bk1)
    os.makedirs(os.path.dirname(bk2), exist_ok=True)
    shutil.copyfile(src, bk2)
    
    execute_reposition_and_ch6(src)
