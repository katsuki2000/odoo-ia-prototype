#!/usr/bin/env python3
import zipfile
import os
import shutil

src_docx = "todo/MASTER_ENI.docx"
test_docx = "/tmp/test_images.docx"
shutil.copyfile(src_docx, test_docx)

replacements = {
    "word/media/image70.png": ".scratch/diagrams_png/fig5_1_usecase.png",
    "word/media/image68.png": ".scratch/diagrams_png/fig5_2_domain.png",
    "word/media/image74.png": ".scratch/diagrams_png/fig5_3_sprint1.png",
    "word/media/image76.png": ".scratch/diagrams_png/fig5_4_sprint3.png",
    "word/media/image72.png": ".scratch/diagrams_png/fig5_5_sprint4.png",
    "word/media/image75.png": ".scratch/diagrams_png/fig5_6_sprint5.png",
    "word/media/image71.png": ".scratch/diagrams_png/fig5_7_sprint7.png",
    "word/media/image77.png": ".scratch/diagrams_png/fig5_8_sprint9.png",
    "word/media/image80.png": ".scratch/diagrams_png/fig5_9_global_classes.png",
    "word/media/image79.png": ".scratch/diagrams_png/fig5_10_package.png",
    "word/media/image78.png": ".scratch/diagrams_png/fig5_11_deployment.png",
}

# Update zip archive
temp_zip = "/tmp/temp_replaced.zip"
with zipfile.ZipFile(test_docx, 'r') as zin, zipfile.ZipFile(temp_zip, 'w') as zout:
    for item in zin.infolist():
        if item.filename in replacements:
            rep_path = replacements[item.filename]
            with open(rep_path, 'rb') as f:
                zout.writestr(item, f.read())
            print(f"Replaced {item.filename} with {rep_path}")
        else:
            zout.writestr(item, zin.read(item.filename))

os.replace(temp_zip, test_docx)
print("Updated /tmp/test_images.docx successfully!")

import docx
doc = docx.Document(test_docx)
print(f"Verified reloaded docx: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")
