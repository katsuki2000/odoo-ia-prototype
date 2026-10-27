import copy
from docx import Document

SRC = 'todo/MASTER_ENI.docx'
TEXTS = [
    "Ce mémoire a porté sur la conception et la réalisation d'un assistant conversationnel privé intégré à Odoo 17, au sein de "
    "l'entreprise MIRO K.E. L'objectif était de permettre aux utilisateurs métier d'obtenir des réponses compréhensibles, "
    "sourcées et conformes à leurs droits d'accès, sans confier les données de l'entreprise à un fournisseur d'intelligence "
    "artificielle externe.",
    "Le système réalisé repose sur deux composants. Le module Odoo miro_ai fournit l'interface de conversation dans la barre "
    "supérieure d'Odoo, signe chaque requête par HMAC, relaie l'identité de l'utilisateur et sa société, suit la tâche de façon "
    "asynchrone et journalise les échanges. Le backend FastAPI, orchestré par LangGraph, choisit la source de la réponse : un "
    "moteur de recherche documentaire hybride, qui fusionne une recherche vectorielle (bge-m3) et une recherche lexicale (BM25) "
    "par Reciprocal Rank Fusion et cite ses sources, ou des outils typés de lecture des données d'Odoo, en lecture seule et sous "
    "l'identité de l'utilisateur. La génération est confiée à un modèle ouvert, Qwen3 14B, servi par un serveur d'inférence "
    "privé : une instance AWS en déploiement et un hôte Kaggle temporaire pour les essais.",
    "Les vérifications effectuées montrent un système fonctionnel de bout en bout : la suite de tests automatisés compte plus "
    "d'un millier de tests réussis, et une série de dix questions métier a obtenu une réponse avec deux profils d'utilisateurs. "
    "Elles révèlent aussi des limites que nous assumons : un temps de réponse de l'ordre de la minute sur le chemin d'analyse du "
    "système, une évaluation de la justesse métier encore partielle et des tests d'intégration qui dépendent d'une instance "
    "Odoo préparée. Un test logiciel réussi ne prouve pas à lui seul la précision métier des réponses ; celle-ci doit être "
    "confirmée par une évaluation reproductible sur des cas définis.",
    "Plusieurs perspectives prolongent ce travail. La principale est l'agent Code Intelligence, présenté au Chapitre 10 : en "
    "construisant un graphe du code et de la configuration de l'instance, il vise à fonder chaque réponse sur des preuves "
    "consultées et à signaler ce qu'il ne peut pas établir. Il a été validé sur des cas réels de validation, d'états de "
    "processus et de surcharge de méthodes, mais sa résolution des termes métier et son refus en l'absence de preuve restent à "
    "fiabiliser. S'y ajoutent le serveur MCP dédié à Odoo évoqué au Chapitre 3, l'extension aux versions plus anciennes d'Odoo "
    "rencontrées chez d'autres clients et la rédaction d'un protocole d'évaluation figé.",
    "Sur le plan personnel, ce travail m'a permis de mettre en pratique, sur un cas concret, les connaissances acquises à "
    "l'École Nationale d'Informatique, et d'approfondir le développement de modules Odoo, la construction d'un service "
    "d'intelligence artificielle privé et la rigueur qui consiste à distinguer ce qu'un système établit de ce qu'il suppose.",
]

doc = Document(SRC)
start = next(i for i, p in enumerate(doc.paragraphs) if p.style.name == 'Heading 1' and p.text.strip() == 'CONCLUSION')
old = []
for p in doc.paragraphs[start + 1:]:
    if p.style.name == 'Heading 1':
        break
    if p.text.strip():
        old.append(p)
assert len(old) == 4 and old[0].text.startswith('Ce stage')
old[0].runs[0].text = TEXTS[0]
for run in old[0].runs[1:]:
    run._r.getparent().remove(run._r)
previous = old[0]
for text in TEXTS[1:]:
    target = old.pop(1) if len(old) > 1 else None
    if target is None:
        clone = copy.deepcopy(old[0]._p)
        previous._p.addnext(clone)
        target = next(p for p in doc.paragraphs if p._p is clone)
    target.runs[0].text = text
    for run in target.runs[1:]:
        run._r.getparent().remove(run._r)
    previous = target
doc.save(SRC)
print('conclusion written')
