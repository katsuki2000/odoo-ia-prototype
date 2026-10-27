#!/usr/bin/env python3
"""
Test repositioning Chapter 5 with exact index 501.
"""
import shutil
import docx

def test_reposition():
    shutil.copyfile('todo/MASTER_ENI.docx', '/tmp/test_MASTER_ENI.docx')
    doc = docx.Document('/tmp/test_MASTER_ENI.docx')
    body = doc._body._element

    old_ch5_h = None
    new_ch5_h = None
    conclusion_h = None

    for idx, child in enumerate(body):
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            t = p.text.strip().upper()
            if 'CHAPITRE 5' in t and p.style.name == 'Heading 2':
                if old_ch5_h is None:
                    old_ch5_h = (idx, child, p.text)
                elif new_ch5_h is None:
                    new_ch5_h = (idx, child, p.text)
            elif 'CONCLUSION' in t and p.style.name == 'Heading 1' and conclusion_h is None:
                conclusion_h = (idx, child, p.text)

    print(f"Old Ch5 at {old_ch5_h[0]}: {old_ch5_h[2]}")
    print(f"New Ch5 at {new_ch5_h[0]}: {new_ch5_h[2]}")
    print(f"Conclusion at {conclusion_h[0]}: {conclusion_h[2]}")

    assert old_ch5_h[0] == 501, f"Expected 501, got {old_ch5_h[0]}"
    assert new_ch5_h[0] == 1083, f"Expected 1083, got {new_ch5_h[0]}"
    assert conclusion_h[0] == 1281, f"Expected 1281, got {conclusion_h[0]}"

    # Elements of new Chapter 5: from new_ch5_h to conclusion_h
    new_ch5_elements = []
    curr = new_ch5_h[1]
    while curr is not None and curr != conclusion_h[1]:
        new_ch5_elements.append(curr)
        curr = curr.getnext()

    print(f"New Chapter 5 elements count: {len(new_ch5_elements)}")

    # Legacy elements: from old_ch5_h to new_ch5_h
    legacy_elements = []
    curr = old_ch5_h[1]
    while curr is not None and curr != new_ch5_h[1]:
        legacy_elements.append(curr)
        curr = curr.getnext()

    print(f"Legacy elements count: {len(legacy_elements)}")

    # Move new_ch5_elements before old_ch5_h[1]
    for elem in new_ch5_elements:
        old_ch5_h[1].addprevious(elem)

    # Delete legacy_elements
    for elem in legacy_elements:
        elem.getparent().remove(elem)

    doc.save('/tmp/test_MASTER_ENI.docx')
    print("Test document saved successfully!")

    # Verify
    doc_reloaded = docx.Document('/tmp/test_MASTER_ENI.docx')
    print("Reloaded! Paragraphs:", len(doc_reloaded.paragraphs), "Tables:", len(doc_reloaded.tables))

    print("\nAll Headings in reloaded document:")
    for i, p in enumerate(doc_reloaded.paragraphs):
        if p.style.name.startswith('Heading'):
            print(f"  {i:04d} [{p.style.name}]: {p.text.strip()[:60]}")

if __name__ == '__main__':
    test_reposition()
