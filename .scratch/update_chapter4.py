#!/usr/bin/env python3
"""
Script to update Chapitre 4 (Analyse préalable) in MASTER_ENI.docx directly.
Replaces legacy Bachelor Brevo/apps menu Chapter 4 with Master-level academic content:
- 4.1. Analyse de l'existant (organisation, processus métier, inventaire matériels/logiciels)
- 4.2. Critique de l'existant (points forts, points faibles, risques de Shadow AI et fuite de données)
- 4.3. Conception avant-projet :
  - 4.3.1. Proposition des solutions (Cloud SaaS vs Statu quo vs IA privée intégrée miro_ai)
  - 4.3.2. Analyse comparative et solution retenue (Tableau 11)
  - 4.3.3. Méthodes de conception et outils retenus (SCRUM vs Merise, UML, tableaux 12 à 15)

Guarantees:
- Everything before child 439 (Cover, Intro, Ch1 ENI, Ch2 Miro, Ch3 Description) is 100% INTACT.
- Everything from child 554 (Chapitre 5. Analyse conceptuelle) onwards is 100% INTACT.
"""

import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_run_font(run, name='Times New Roman', size_pt=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def create_styled_paragraph(doc, text="", style='Normal', bold=False, italic=False, size_pt=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.style = style
    p.paragraph_format.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_run_font(r, size_pt=size_pt, bold=bold, italic=italic)
    return p

def create_subheading(doc, text, level=3, size_pt=14):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=size_pt, bold=True)
    return p

def create_bullet(doc, text, prefix="", size_pt=12):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_bullet = p.add_run("• ")
    set_run_font(r_bullet, size_pt=size_pt, bold=True)
    
    if prefix:
        r_pre = p.add_run(prefix)
        set_run_font(r_pre, size_pt=size_pt, bold=True)
        
    r_text = p.add_run(text)
    set_run_font(r_text, size_pt=size_pt)
    return p

def create_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.style = 'Table10'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for idx, title in enumerate(headers):
        hdr_cells[idx].text = title
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            set_run_font(r, size_pt=10.5, bold=True)
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0 and len(str(cell_value)) <= 8:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 1 and len(str(cell_value)) <= 6:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                set_run_font(r, size_pt=10)
                
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    return tbl

def update_chapter_4(doc_path):
    print(f"Loading document: {doc_path}...")
    doc = docx.Document(doc_path)
    body = doc._body._element

    chap4_h = None
    chap5_h = None

    for idx, child in enumerate(body):
        tag = child.tag.split('}')[-1]
        if tag == 'p' and idx >= 400:
            p = docx.text.paragraph.Paragraph(child, doc)
            t = p.text.strip().upper()
            if 'CHAPITRE 4' in t and 'ANALYSE' in t and chap4_h is None:
                chap4_h = (idx, child, p.text)
            elif 'CHAPITRE 5' in t and 'ANALYSE' in t and chap5_h is None:
                chap5_h = (idx, child, p.text)

    assert chap4_h is not None, "Chapitre 4 heading not found!"
    assert chap5_h is not None, "Chapitre 5 heading not found!"
    assert chap4_h[0] < chap5_h[0], "Chapitre 4 and 5 out of order!"

    print(f"Chapitre 4 heading at index {chap4_h[0]}: '{chap4_h[2]}'")
    print(f"Chapitre 5 heading at index {chap5_h[0]}: '{chap5_h[2]}'")

    # Delete all elements strictly between chap4_h and chap5_h
    elements_to_delete = []
    curr = chap4_h[1].getnext()
    while curr is not None and curr != chap5_h[1]:
        elements_to_delete.append(curr)
        curr = curr.getnext()

    print(f"Removing {len(elements_to_delete)} legacy elements in Chapter 4...")
    for elem in elements_to_delete:
        elem.getparent().remove(elem)

    # Now insert Chapter 4 elements before chap5_h[1]
    ref = chap5_h[1]

    # Introduction Chapitre 4
    p_intro = create_styled_paragraph(doc, "L'intégration réussie d'une solution d'intelligence artificielle au sein d'un Progiciel de Gestion Intégré exige une analyse rigoureuse des processus organisationnels existants et des flux informationnels. Ce chapitre présente l'analyse préalable du projet en trois étapes majeures : l'étude de l'existant au sein de l'entreprise d'accueil MIRO K.E. et de son écosystème d'utilisateurs Odoo 17, la critique méthodique identifiant les verrous technologiques et les risques de sécurité (notamment l'usage non maîtrisé d'IA externes), et enfin la conception avant-projet détaillant la formulation des architectures envisageables, la justification de la solution retenue ainsi que le choix des démarches méthodologiques et des outils de modélisation.", space_before=6)
    ref.addprevious(p_intro._p)

    # 4.1 Analyse de l'existant
    h4_1 = create_subheading(doc, "4.1. Analyse de l’existant", size_pt=14)
    ref.addprevious(h4_1._p)

    h4_1_1 = create_subheading(doc, "4.1.1. Organisation actuelle et processus métier", size_pt=12)
    ref.addprevious(h4_1_1._p)

    p4_1_1_1 = create_styled_paragraph(doc, "L'entreprise MIRO K.E. exploite la plateforme Odoo 17 pour piloter ses activités internes et déploie cette même solution auprès de ses clients d'horizons variés (négoce, services, industrie, distribution). Au sein de cette organisation, le Progiciel de Gestion Intégré concentre les opérations critiques quotidiennes à travers plusieurs modules métiers interconnectés : gestion commerciale et devis, facturation et comptabilité analytique, gestion des stocks et réapprovisionnements, gestion des relations clients (CRM) et assistance technique.")
    ref.addprevious(p4_1_1_1._p)

    p4_1_1_2 = create_styled_paragraph(doc, "L'analyse opérationnelle révèle deux flux de traitement informationnel majeurs, aujourd'hui exécutés de manière cloisonnée et manuelle :")
    ref.addprevious(p4_1_1_2._p)

    flux_list = [
        ("La recherche documentaire et procédurale : ", "Les collaborateurs (consultants d'intégration, comptables, gestionnaires administratifs) doivent fréquemment consulter les guides méthodologiques, les fiches techniques d'implémentation et les règles de gestion spécifiques aux configurations déployées. Ces ressources documentaires sont réparties sous divers formats statiques (manuels PDF, notes de versions, dépôts de documentation Markdown et wikis internes). La recherche d'une procédure précise requiert une navigation manuelle chronophage et se révèle souvent infructueuse face à l'abondance des référentiels."),
        ("L'interrogation des données transactionnelles de l'ERP : ", "Lorsqu'un responsable commercial ou un gestionnaire de stock souhaite obtenir un indicateur de gestion (par exemple, le montant cumulé des commandes en cours d'un client, la liste des factures impayées échues ou la disponibilité immédiate d'un article multi-sociétés), il est contraint d'ouvrir le menu dédié, de manipuler des filtres de recherche complexes, de basculer entre plusieurs vues (liste, formulaire, pivot) ou d'exporter les enregistrements sous tableur externe afin d'opérer des regroupements manuels."),
    ]
    for pre, fl in flux_list:
        p_fl = create_bullet(doc, fl, prefix=pre)
        ref.addprevious(p_fl._p)

    p4_1_1_3 = create_styled_paragraph(doc, "Sur le plan organisationnel, les catégories de personnel impliquées dans ces traitements regroupent la Direction technique et commerciale, les consultants intégrateurs Odoo, les gestionnaires de paie et de facturation, ainsi que les équipes de support client de Miro K.E.")
    ref.addprevious(p4_1_1_3._p)

    h4_1_2 = create_subheading(doc, "4.1.2. Inventaire des moyens matériels et logiciels", size_pt=12)
    ref.addprevious(h4_1_2._p)

    p4_1_2_1 = create_styled_paragraph(doc, "L'infrastructure technique actuellement en exploitation au sein de Miro K.E. pour supporter les environnements Odoo 17 est synthétisée dans le Tableau 10.")
    ref.addprevious(p4_1_2_1._p)

    # Caption Table 10
    p_cap10 = create_styled_paragraph(doc, "Tableau 10. Inventaire des infrastructures matérielles et logicielles existantes chez Miro K.E.", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap10._p)

    headers_tbl10 = ["Catégorie", "Composant / Équipement", "Spécifications techniques", "Rôle opérationnel"]
    data_tbl10 = [
        ["Matériel", "Serveurs d'hébergement Odoo", "Processeurs multi-cœurs x86_64, 32 Go RAM, SSD NVMe", "Hébergement des instances de test, préproduction et production."],
        ["Matériel", "Postes de travail collaborateurs", "PC portables et fixes Intel Core i5/i7, 16 Go RAM", "Postes clients de consultation, de gestion administrative et de développement."],
        ["Logiciel", "Système d'exploitation serveur", "Linux Debian 12 / Ubuntu Server 22.04 LTS", "Environnement sécurisé d'exécution des services d'entreprise."],
        ["Logiciel", "Progiciel de Gestion Intégré", "Odoo 17.0 (Community et Enterprise)", "Plateforme centrale de gestion d'entreprise (CRM, Vente, Stock, Compta)."],
        ["Logiciel", "Système de Gestion de BDD", "PostgreSQL 14 / 16", "Stockage persistant des données relationnelles et transactionnelles."],
        ["Logiciel", "Navigateurs Web clients", "Google Chrome, Mozilla Firefox (dernières versions)", "Accès web des utilisateurs aux interfaces de l'ERP."],
        ["Logiciel", "Outils d'assistance existants", "Recherche textuelle native d'Odoo, tableurs", "Recherche élémentaire par sous-chaînes de caractères sans compréhension sémantique."],
    ]
    tbl10 = create_styled_table(doc, headers_tbl10, data_tbl10, col_widths=[1.1, 1.8, 2.2, 2.0])
    ref.addprevious(tbl10._tbl)

    p_sp1 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp1._p)

    # 4.2 Critique de l'existant
    h4_2 = create_subheading(doc, "4.2. Critique de l’existant", size_pt=14)
    ref.addprevious(h4_2._p)

    p4_2_intro = create_styled_paragraph(doc, "L'évaluation approfondie du fonctionnement actuel met en lumière des atouts structurels indéniables, mais révèle également des goulots d'étranglement critiques ainsi que des risques de sécurité majeurs.")
    ref.addprevious(p4_2_intro._p)

    h4_2_pf = create_subheading(doc, "Points forts de l'organisation actuelle :", size_pt=12)
    ref.addprevious(h4_2_pf._p)

    pf_list = [
        ("Centralisation et cohérence des données : ", "L'utilisation d'Odoo 17 élimine la redondance des enregistrements en regroupant clients, commandes, stocks et factures dans une base unique et normalisée."),
        ("Sécurité applicative native robuste : ", "Odoo dispose d'un système de contrôle d'accès éprouvé (RBAC), combinant autorisations de modèles (ir.model.access) et filtres dynamiques au niveau des lignes (ir.rule), garantissant un cloisonnement strict entre utilisateurs et sociétés."),
        ("Architecture web moderne : ", "Le framework JavaScript OWL 2.0 offre une interface utilisateur réactive, composable et modulaire adaptée aux environnements web d'entreprise."),
    ]
    for pre, pf in pf_list:
        p_pf = create_bullet(doc, pf, prefix=pre)
        ref.addprevious(p_pf._p)

    h4_2_lim = create_subheading(doc, "Points faibles et vulnérabilités identifiés :", size_pt=12)
    ref.addprevious(h4_2_lim._p)

    lim_list = [
        ("Surcharge cognitive et perte de productivité : ", "L'arborescence étendue des menus d'Odoo 17 rend complexe l'accès rapide aux informations synthétiques pour les utilisateurs non spécialistes, générant des lenteurs et des interruptions répétées du travail opérationnel."),
        ("Incapacité de la recherche native à traiter le langage naturel : ", "La recherche textuelle standard d'Odoo s'effectue par correspondance stricte de chaînes (opérateur ILIKE SQL). Elle s'avère totalement inopérante pour interpréter des questions conceptuelles, des formulations synonymes ou des requêtes métier complexes."),
        ("Fragmentation des connaissances documentaires : ", "Les documents techniques et procéduraux sont dissociés des flux opérationnels de l'ERP, privant les utilisateurs d'une assistance contextuelle immédiate lors de l'exécution de leurs tâches."),
        ("Risque critique de fuite de données et « Shadow AI » : ", "Face au besoin quotidien de synthétiser des textes ou d'analyser des situations complexes, certains collaborateurs ont recours à des assistants IA grand public en mode cloud (ChatGPT, Claude, etc.) en y injectant des extraits de données financières, clients ou RH. Cette pratique expose l'entreprise à des violations graves du secret commercial et des obligations légales de protection des données (loi n° 2014-038 à Madagascar et RGPD)."),
        ("Absence d'étanchéité des privilèges dans les outils externes : ", "Aucun outil conversationnel généraliste externe n'est capable de respecter le contexte de session, l'identifiant strict de l'utilisateur (uid) et les règles multi-sociétés d'Odoo, ce qui interdit tout interfaçage direct sans supervision de sécurité."),
    ]
    for pre, lim in lim_list:
        p_lim = create_bullet(doc, lim, prefix=pre)
        ref.addprevious(p_lim._p)

    # 4.3 Conception avant-projet
    h4_3 = create_subheading(doc, "4.3. Conception avant-projet", size_pt=14)
    ref.addprevious(h4_3._p)

    h4_3_1 = create_subheading(doc, "4.3.1. Proposition des solutions", size_pt=12)
    ref.addprevious(h4_3_1._p)

    p4_3_1_intro = create_styled_paragraph(doc, "Pour répondre aux limites identifiées et doter l'entreprise d'une assistance intelligente sécurisée, trois orientations architecturales ont été analysées :")
    ref.addprevious(p4_3_1_intro._p)

    sol_list = [
        ("Solution 1 : Intégration d'un service d'IA SaaS sur Cloud public (ex. API OpenAI GPT-4 / Claude) : ", "Cette approche consiste à développer un connecteur Odoo expédiant les questions des utilisateurs et des extraits de la base de données vers des serveurs hébergés par des fournisseurs tiers. Si elle bénéficie d'une puissance de calcul immédiate sans investissement matériel local, elle présente des risques rédhibitoires : fuite de données d'entreprise sensibles hors du périmètre souverain, coûts récurrents élevés indexés sur la consommation de jetons (tokens), dépendance absolue vis-à-vis d'une connexion Internet stable et impossibilité de garantir l'absence d'élévation de privilèges au niveau de l'ORM Odoo."),
        ("Solution 2 : Maintien du statu quo avec optimisation des outils manuels existants : ", "Cette solution préconise de former davantage les utilisateurs à l'ergonomie standard d'Odoo et de développer des tableaux de bord statiques supplémentaires. Bien qu'elle n'engendre aucun investissement matériel, elle ne résout ni la fracture cognitive liée à la multiplicité des menus, ni la fragmentation documentaire, et n'endigue en rien le recours informel au « Shadow AI »."),
        ("Solution 3 : Conception et déploiement d'une plateforme d'IA privée souveraine intégrée (Addon miro_ai + Backend FastAPI/LangGraph + LLM local + RAG hybride + ORM sécurisé) : ", "Cette architecture déploie une solution d'intelligence artificielle 100 % hébergée sur l'infrastructure privée de Miro K.E. Elle intègre un module addon Odoo 17 (miro_ai) conçu en JavaScript OWL pour l'IHM et des contrôleurs asynchrones HMAC, couplé à un middleware d'orchestration FastAPI/LangGraph opérant des modèles open-weight locaux (Ollama/vLLM, BGE-M3). Le moteur RAG hybride résout l'accès à la documentation avec citation des sources, tandis que l'agent ORM interroge l'ERP sous l'identité stricte du demandeur (uid) en respectant rigoureusement les droits RBAC sans jamais exécuter de sudo()."),
    ]
    for pre, sol in sol_list:
        p_sol = create_bullet(doc, sol, prefix=pre)
        ref.addprevious(p_sol._p)

    h4_3_2 = create_subheading(doc, "4.3.2. Analyse comparative et solution retenue", size_pt=12)
    ref.addprevious(h4_3_2._p)

    p_comp_intro = create_styled_paragraph(doc, "Le Tableau 11 présente une analyse comparative multicritères des trois solutions envisagées afin d'éclairer la décision stratégique.")
    ref.addprevious(p_comp_intro._p)

    # Caption Table 11
    p_cap11 = create_styled_paragraph(doc, "Tableau 11. Analyse comparative des solutions architecturales envisagées", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap11._p)

    headers_tbl11 = ["Critères d'évaluation", "Solution 1 : IA Cloud SaaS", "Solution 2 : Statu Quo", "Solution 3 : IA Privée Souveraine (miro_ai)"]
    data_tbl11 = [
        ["Souveraineté et confidentialité des données", "Faible (données transmises à des tiers)", "Moyenne (données internes mais risque Shadow AI)", "Maximale (traitement 100% on-premise en circuit fermé)"],
        ["Respect du contrôle d'accès Odoo (RBAC)", "Complexe et risqué (filtrage difficile en amont)", "Natif (menus Odoo standard)", "Strict (héritage dynamique de l'uid et des ir.rule sans sudo)"],
        ["Modèle économique et coûts récurrents", "Coûteux (abonnement et facturation aux tokens)", "Nul (aucun coût d'IA additionnel)", "Maîtrisé (amortissement du serveur matériel, zéro coût token)"],
        ["Ergonomie et intégration dans l'ERP", "Moyenne (widget déconnecté ou externe)", "Faible (navigation complexe manuelle)", "Excellente (composant natif OWL dans le systray d'Odoo 17)"],
        ["Résistance aux hallucinations métier", "Faible (modèle généraliste sans ancrage)", "Non applicable", "Élevée (RAG hybride avec citations et validation ORM)"],
        ["Indépendance réseau et disponibilité", "Dépendance critique à la bande passante", "Totale en réseau local", "Totale (opère même en cas de coupure Internet)"],
    ]
    tbl11 = create_styled_table(doc, headers_tbl11, data_tbl11, col_widths=[1.8, 1.8, 1.5, 2.0])
    ref.addprevious(tbl11._tbl)

    p_sp2 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp2._p)

    p_retenue = create_styled_paragraph(doc, "Au vu de cette analyse multicritères, la Solution 3 est formellement retenue. Elle est la seule à concilier les exigences absolues de souveraineté des données, de respect strict de la politique de sécurité RBAC d'Odoo, d'absence totale de coûts récurrents à l'usage et d'intégration fluide au poste de travail des utilisateurs grâce à l'addon miro_ai.")
    ref.addprevious(p_retenue._p)

    # 4.3.3 Méthodes de conception et outils retenus
    h4_3_3 = create_subheading(doc, "4.3.3. Méthodes de conception et outils retenus", size_pt=12)
    ref.addprevious(h4_3_3._p)

    p_meth_intro = create_styled_paragraph(doc, "La réussite d'un projet d'ingénierie logicielle couplant un ERP d'entreprise et des composants d'intelligence artificielle requiert une démarche méthodologique rigoureuse et des choix technologiques ciblés.")
    ref.addprevious(p_meth_intro._p)

    h_scrum = create_subheading(doc, "Démarche de conduite de projet : Processus Agile SCRUM", size_pt=11)
    ref.addprevious(h_scrum._p)

    p_scrum = create_styled_paragraph(doc, "Contrairement aux cycles séquentiels rigides (tels que le cycle en V ou la méthode MERISE) qui fixent l'intégralité des spécifications avant le début du développement, le développement de systèmes à base d'agents IA et de RAG requiert une approche hautement itérative. Les réglages des modèles d'embedding, le seuillage de pertinence et la conception des prompts nécessitent des validations expérimentales rapprochées. C'est pourquoi la méthodologie Agile SCRUM a été adoptée, organisée en sprints de deux semaines jalonnés de démonstrations fonctionnelles auprès de l'équipe technique de Miro K.E. Le Tableau 12 confronte les caractéristiques de ces deux approches.")
    ref.addprevious(p_scrum._p)

    # Caption Table 12
    p_cap12 = create_styled_paragraph(doc, "Tableau 12. Analyse comparative de la méthode Agile SCRUM et du cycle séquentiel (MERISE / Cascade)", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap12._p)

    headers_tbl12 = ["Dimension méthodologique", "Cycle séquentiel (MERISE / Cascade)", "Démarche Agile SCRUM (Retenue)"]
    data_tbl12 = [
        ["Gestion des exigences", "Spécifications exhaustives figées dès le démarrage", "Product Backlog dynamique et priorisé selon la valeur métier"],
        ["Adaptabilité aux aléas de l'IA", "Rigide face aux imprévus de performance des modèles", "Forte réactivité grâce aux revues de sprint régulières"],
        ["Visibilité des livrables", "Validation tardive lors de la phase de recette finale", "Incréments logiciels testables livrés à chaque sprint"],
        ["Implication des parties prenantes", "Ponctuelle (au début et à la livraison finale)", "Continue (collaboration étroite avec les encadreurs)"],
    ]
    tbl12 = create_styled_table(doc, headers_tbl12, data_tbl12, col_widths=[2.0, 2.5, 2.6])
    ref.addprevious(tbl12._tbl)

    p_sp3 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp3._p)

    h_uml = create_subheading(doc, "Langage de modélisation retenu : UML 2.5", size_pt=11)
    ref.addprevious(h_uml._p)

    p_uml = create_styled_paragraph(doc, "Le langage de modélisation unifié UML (Unified Modeling Language), standard international de l'OMG, a été choisi pour concevoir et documenter les différentes facettes du système. Dans le cadre de ce projet, les diagrammes suivants seront particulièrement exploités :")
    ref.addprevious(p_uml._p)

    uml_list = [
        ("Le diagramme de cas d'utilisation : ", "pour formaliser les interactions entre les utilisateurs finaux, les administrateurs et l'assistant IA miro_ai."),
        ("Les diagrammes de séquence : ", "pour expliciter le flux temporel asynchrone des requêtes, depuis le clic de l'utilisateur dans l'IHM OWL jusqu'à la consultation de l'ORM Odoo via FastAPI et la restitution finale."),
        ("Le diagramme de classes et de paquetages : ", "pour structurer l'architecture modulaire du backend (nœuds LangGraph, modèles de configuration Odoo, contrôleurs de sécurité)."),
        ("Le diagramme de déploiement : ", "pour illustrer la topologie réseau, l'isolation des conteneurs Docker et les liaisons inter-processus sécurisées par HMAC."),
    ]
    for pre, um in uml_list:
        p_um = create_bullet(doc, um, prefix=pre)
        ref.addprevious(p_um._p)

    h_tech = create_subheading(doc, "Analyses comparatives des composants technologiques clés", size_pt=11)
    ref.addprevious(h_tech._p)

    p_tech_intro = create_styled_paragraph(doc, "Pour chaque couche critique de l'architecture, une sélection rigoureuse d'outils et de bibliothèques a été opérée, comme le détaillent les tableaux comparatifs ci-après :")
    ref.addprevious(p_tech_intro._p)

    # Table 13: Moteurs LLM
    p_cap13 = create_styled_paragraph(doc, "Tableau 13. Analyse comparative des moteurs d'inférence LLM locaux", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap13._p)

    headers_tbl13 = ["Moteur d'inférence", "Points forts", "Limites", "Décision et usage"]
    data_tbl13 = [
        ["Ollama", "Déploiement simple, gestion unifiée des modèles GGUF, excellente compatibilité", "Optimisation multi-requêtes simultanées limitée", "Retenu pour le prototypage rapide et l'environnement de développement"],
        ["vLLM", "Haut débit, gestion avancée de la mémoire (PagedAttention), idéal pour production", "Nécessite des GPU dédiés avec environnement CUDA natif", "Retenu pour le déploiement serveur de production à forte charge"],
        ["HuggingFace TGI", "Excellente intégration à l'écosystème HF, support de la quantification", "Configuration complexe et empreinte mémoire importante", "Écarté au profit d'Ollama et vLLM plus modulaires"],
    ]
    tbl13 = create_styled_table(doc, headers_tbl13, data_tbl13, col_widths=[1.5, 2.2, 1.8, 1.8])
    ref.addprevious(tbl13._tbl)

    p_sp4 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp4._p)

    # Table 14: RAG
    p_cap14 = create_styled_paragraph(doc, "Tableau 14. Analyse comparative des stratégies de recherche documentaire pour le RAG", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap14._p)

    headers_tbl14 = ["Approche RAG", "Comportement sur le vocabulaire métier", "Comportement conceptuel", "Décision et justification"]
    data_tbl14 = [
        ["Vectoriel dense pur (Dense-only)", "Faible : peine sur les codes d'articles et termes techniques rares", "Excellent : capte les synonymes et analogies sémantiques", "Insuffisant pour la précision chirurgicale requise par un ERP"],
        ["Lexical pur (BM25 seul)", "Excellent sur les identifiants exacts et codes comptables", "Nul : incapable d'interpréter des reformulations ou périphrases", "Insuffisant face aux formulations libres des utilisateurs"],
        ["RAG Hybride Dense + BM25 avec fusion RRF (Retenu)", "Maximale : réconciliation des codes précis et des questions sémantiques", "Optimale : combinaison robuste des scores avec dédoublonnage", "Retenu : garantit un rappel supérieur à 90% sur la documentation Odoo"],
    ]
    tbl14 = create_styled_table(doc, headers_tbl14, data_tbl14, col_widths=[1.8, 1.8, 1.8, 1.9])
    ref.addprevious(tbl14._tbl)

    p_sp5 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp5._p)

    # Table 15: Frontend
    p_cap15 = create_styled_paragraph(doc, "Tableau 15. Analyse comparative des modes d'intégration de l'interface utilisateur dans Odoo 17", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    ref.addprevious(p_cap15._p)

    headers_tbl15 = ["Mode d'intégration IHM", "Expérience utilisateur (UX)", "Sécurité et contexte de session", "Décision et justification"]
    data_tbl15 = [
        ["Iframe web externe", "Médiocre : rupture visuelle et redimensionnement malaisé", "Complexe : risque d'usurpation de session par cross-frame", "Écarté pour des motifs d'ergonomie et de sécurité"],
        ["Application web indépendante", "Faible : contraint l'utilisateur à changer d'onglet de travail", "Nécessite une ré-authentification séparée", "Écarté car incompatible avec l'usage opérationnel fluide au sein d'Odoo"],
        ["Composant natif OWL 2.0 (Retenu dans miro_ai)", "Excellente : intégration systray et volet rétractable fluide", "Natif : transmission directe et vérifiée du contexte utilisateur", "Retenu : respect absolu des standards technologiques d'Odoo 17"],
    ]
    tbl15 = create_styled_table(doc, headers_tbl15, data_tbl15, col_widths=[1.8, 1.8, 1.8, 1.9])
    ref.addprevious(tbl15._tbl)

    p_sp6 = create_styled_paragraph(doc, "")
    ref.addprevious(p_sp6._p)

    print("Chapter 4 successfully inserted.")

    print(f"Saving document: {doc_path}...")
    doc.save(doc_path)
    print("Document successfully saved!")

if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else '/home/user/Downloads/MASTER_ENI.docx'
    update_chapter_4(target)
