import copy
from docx import Document

SRC = 'todo/MASTER_ENI.docx'
doc = Document(SRC)
paragraphs = doc.paragraphs


def replace_once(old, new):
    hits = [(p, r) for p in paragraphs for r in p.runs if old in r.text]
    assert len(hits) == 1, f'{len(hits)} match(es) for: {old[:60]}'
    run = hits[0][1]
    run.text = run.text.replace(old, new)


def set_cell(cell, text):
    paragraph = cell.paragraphs[0]
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)


def find_table(first_header, first_cell):
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip() == first_header and table.rows[1].cells[0].text.strip().startswith(first_cell):
            return table
    raise AssertionError(first_cell)


def find_row(table, label):
    for row in table.rows:
        if row.cells[0].text.strip().startswith(label):
            return row
    raise AssertionError(label)


replace_once('chunks de 512 tokens avec un recouvrement de 64 tokens', 'chunks de 512 unités avec un recouvrement de 50 unités')
replace_once('(ChromaDB / Qdrant)', '(ChromaDB)')
replace_once("opérant en circuit fermé sur les serveurs de Miro K.E.,",
             "opérant sur une infrastructure privée (instance AWS dédiée en déploiement, hôte Kaggle temporaire pendant le développement),")
replace_once(
    "Le modèle Qwen 2.5 (14B) quantifié en format GGUF Q4_K_M est retenu comme moteur génératif de référence. Il combine une capacité de raisonnement et de fidélité aux instructions remarquable, un excellent traitement de la langue française et une exécution fluide sur notre serveur équipé d'une carte GPU NVIDIA RTX (tenant dans 12 Go de VRAM). Pour les environnements de test plus contraints, Llama 3.1 (8B) sert de solution alternative hautement réactive.",
    "Le modèle Qwen3 (14B) quantifié en format GGUF Q4_K_M est retenu comme moteur génératif de référence. Il combine une bonne capacité de raisonnement et de suivi des instructions et un traitement correct de la langue française ; sur le serveur d'inférence, il est entièrement chargé en mémoire GPU avec une fenêtre de contexte de 32 768 unités. Pour les environnements plus contraints, Llama 3.1 (8B) reste une alternative envisagée, non utilisée dans la configuration actuelle.")
replace_once("En raison du délai incompressible requis par l'inférence locale (de 2 à 5 secondes selon la complexité)",
             "En raison du délai de traitement, de quelques secondes à plus d'une minute selon la complexité de la question (voir la section 9.3)")
replace_once('Confinement réseau absolu : ', "Maîtrise de l'infrastructure d'inférence : ")
replace_once("Les communications s'opèrent exclusivement au sein du réseau local d'entreprise ou du réseau virtuel Docker interne, sans aucune exposition sur Internet.",
             "Aucun fournisseur d'IA externe n'est appelé et il n'existe aucun repli implicite vers un service tiers : l'adresse du serveur d'inférence est fournie par l'environnement et un hôte distant doit être explicitement autorisé. Le serveur d'inférence est une instance AWS privée en déploiement ; un hôte Kaggle exposé par un tunnel ngrok sert de cible temporaire pour les essais.")

llm = find_table('Modèle candidat', 'Qwen 2.5')
set_cell(llm.rows[1].cells[0], 'Qwen3 (14B)')
set_cell(find_row(llm, 'Llama 3.1').cells[5], 'Alternative légère, non utilisée dans la configuration actuelle')

params = find_table('Paramètre', 'Taille des chunks')
set_cell(params.rows[1].cells[1], '512 unités')
set_cell(find_row(params, 'Recouvrement').cells[1], '50 unités')
set_cell(find_row(params, 'Pondération RRF').cells[1], 'k = 60 ; poids 0,6 (dense) et 0,3 (BM25)')
temperature = find_row(params, 'Température')
set_cell(temperature.cells[1], '0,3 (réponses documentaires)')
set_cell(temperature.cells[2], "Génération peu variable pour limiter l'extrapolation ; la température par défaut du client LLM est nulle")
topP = find_row(params, 'Top-p')
topP._tr.getparent().remove(topP._tr)

api = find_table('Méthode & Route', 'POST /api/v1/health')
health = api.rows[1]
set_cell(health.cells[0], 'GET /api/v1/health')
set_cell(health.cells[2], 'Jeton HMAC (compte administrateur)')
sessions = find_row(api, 'POST /api/v1/sessions')
set_cell(sessions.cells[0], 'POST /miro_ai/sessions')
set_cell(sessions.cells[1], 'Web Client OWL -> Odoo')
set_cell(sessions.cells[2], 'limit')
set_cell(sessions.cells[3], "HTTP 200 : liste des conversations récentes de l'utilisateur, lue dans le journal miro.ai.log")

replace_once(
    "L'ensemble tourne sous Linux (Debian) dans un environnement virtuel Python, et l'inférence est confiée à un serveur privé compatible avec l'API OpenAI (Ollama ou vLLM), sans repli vers un fournisseur d'IA externe.",
    "L'ensemble tourne sous Linux (Debian) dans un environnement virtuel Python. L'inférence est confiée à un serveur compatible avec l'API OpenAI (Ollama ou vLLM) dont l'adresse est fournie par l'environnement ; aucun fournisseur d'IA externe n'est appelé et il n'existe aucun repli implicite vers un service tiers.")
intro = next(p for p in paragraphs if p.text.startswith('Le Tableau 37 recense'))
targets = copy.deepcopy(intro._p)
intro._p.addnext(targets)
newPara = next(p for p in doc.paragraphs if p._p is targets)
newPara.runs[0].text = (
    "Deux cibles d'inférence sont prévues, choisies par la variable LLM_TARGET. La cible aws, par défaut, est une instance AWS "
    "privée qui sert de serveur d'inférence de déploiement. La cible kaggle est un hôte de test temporaire : un notebook Kaggle "
    "exposé par un tunnel ngrok, utilisé pour les essais pendant le développement. Cette seconde cible est isolée dans une seule "
    "entrée de la configuration et se retire en supprimant cette entrée et les variables KAGGLE_*. Le serveur contrôlé le "
    "26 octobre 2026 exécute Ollama 0.40.8 et sert les modèles qwen3:14b (quantifié Q4_K_M, entièrement en mémoire GPU, fenêtre "
    "de 32 768 unités) et bge-m3. Un hôte distant doit être explicitement autorisé : par défaut, seules les adresses locales le sont."
)

env = find_table('Composant', 'Système')
set_cell(find_row(env, "Serveur d'inférence").cells[1], 'Ollama, compatible vLLM ; instance AWS (déploiement) ou hôte Kaggle (test)')
set_cell(find_row(env, 'Modèle de génération').cells[1], 'qwen3:14b, quantifié Q4_K_M')
steps = find_table('Étape', '1. Dépendances')
set_cell(find_row(steps, '2. Configuration').cells[1],
         "Copie de .env.example vers .env : cible d'inférence (LLM_TARGET), adresse du serveur, mot de passe du compte de service Odoo")

doc.save(SRC)
print('done')
