#!/usr/bin/env python3
"""
Script to update /home/user/Downloads/MASTER_ENI.docx directly with:
1. Cover page title: "Intitulé : Conception, réalisation et intégration d’une intelligence artificielle privée dans Odoo 17"
2. Introduction Générale: Master-level academic text highlighting miro_ai (Odoo addon) + Private AI backend
3. Chapitre 3: Description du projet (Master-level academic text, tables 6, 7, 8, 9, highlighting miro_ai, and mandatory perspectives MCP & Code Intelligence)

Guarantees:
- Chapitre 1 (ENI) and Chapitre 2 (Miro K.E.) are 100% STRICTLY INTACT.
- Partie II, Partie III and beyond are intact.
"""

import sys
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_run_font(run, name='Times New Roman', size_pt=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def create_styled_paragraph(doc, text="", style='Normal', bold=False, italic=False, size_pt=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.style = style
    p.paragraph_format.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_run_font(r, size_pt=size_pt, bold=bold, italic=italic)
    return p

def create_subheading(doc, text, level=3, size_pt=14):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=size_pt, bold=True)
    return p

def create_bullet(doc, text, prefix="", size_pt=12):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_bullet = p.add_run("• ")
    set_run_font(r_bullet, size_pt=size_pt, bold=True)
    
    if prefix:
        r_pre = p.add_run(prefix)
        set_run_font(r_pre, size_pt=size_pt, bold=True)
        
    r_text = p.add_run(text)
    set_run_font(r_text, size_pt=size_pt)
    return p

def create_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.style = 'Table10'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for idx, title in enumerate(headers):
        hdr_cells[idx].text = title
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            set_run_font(r, size_pt=10.5, bold=True)
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0 and len(str(cell_value)) <= 8:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 1 and len(str(cell_value)) <= 6:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                set_run_font(r, size_pt=10)
                
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    return tbl

def update_document(doc_path):
    print(f"Loading document: {doc_path}...")
    doc = docx.Document(doc_path)
    body = doc._body._element

    # 1. Update Cover Page Title (paragraph index 12)
    p12 = doc.paragraphs[12]
    print(f"Updating p12 from: '{p12.text}'")
    p12.text = ""
    r_label = p12.add_run("Intitulé : ")
    set_run_font(r_label, size_pt=12, bold=True)
    r_title = p12.add_run("Conception, réalisation et intégration d’une intelligence artificielle privée dans Odoo 17")
    set_run_font(r_title, size_pt=12, bold=True)
    p12.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2. Locate boundaries with index >= 200 to avoid matching TOC entries
    intro_h = None
    partie_1 = None
    chap3_h = None
    partie_2 = None

    for idx, child in enumerate(body):
        tag = child.tag.split('}')[-1]
        if tag == 'p' and idx >= 200:
            p = docx.text.paragraph.Paragraph(child, doc)
            t = p.text.strip().upper()
            if 'INTRODUCTION GENERALE' in t and intro_h is None:
                intro_h = (idx, child, p.text)
            elif 'PARTIE I' in t and 'PRESENTATION' in t and partie_1 is None:
                partie_1 = (idx, child, p.text)
            elif 'CHAPITRE 3' in t and 'DESCRIPTION' in t and chap3_h is None:
                chap3_h = (idx, child, p.text)
            elif 'PARTIE II' in t and 'ANALYSE ET CONCEPTION' in t and partie_2 is None:
                partie_2 = (idx, child, p.text)

    assert intro_h is not None, "INTRODUCTION GENERALE heading not found!"
    assert partie_1 is not None, "PARTIE I heading not found!"
    assert chap3_h is not None, "Chapitre 3 heading not found!"
    assert partie_2 is not None, "PARTIE II heading not found!"

    print(f"Intro heading found at index: {intro_h[0]} ('{intro_h[2]}')")
    print(f"Partie I heading found at index: {partie_1[0]} ('{partie_1[2]}')")
    print(f"Chapitre 3 heading found at index: {chap3_h[0]} ('{chap3_h[2]}')")
    print(f"Partie II heading found at index: {partie_2[0]} ('{partie_2[2]}')")

    assert intro_h[0] < partie_1[0] < chap3_h[0] < partie_2[0], "Heading order mismatch!"

    # -------------------------------------------------------------
    # 3. Update INTRODUCTION GENERALE
    # -------------------------------------------------------------
    elements_to_delete_intro = []
    curr = intro_h[1].getnext()
    while curr is not None and curr != partie_1[1]:
        elements_to_delete_intro.append(curr)
        curr = curr.getnext()

    print(f"Removing {len(elements_to_delete_intro)} legacy paragraphs in Introduction...")
    for elem in elements_to_delete_intro:
        elem.getparent().remove(elem)

    intro_paragraphs = [
        "À l'ère de la transformation numérique, les systèmes d'information d'entreprise ont atteint un niveau de maturité où l'optimisation des flux opérationnels repose sur l'intégration continue et l'interopérabilité des données. Au cœur de cette architecture, les Progiciels de Gestion Intégrés (PGI / ERP), à l'instar d'Odoo, fédèrent l'ensemble des processus critiques de l'organisation : gestion de la chaîne logistique, comptabilité et finances, relations clients (CRM), ressources humaines et production industrielle. Parallèlement, l'avènement spectaculaire des modèles de fondation et de l'intelligence artificielle générative (LLM — Large Language Models) redéfinit en profondeur les paradigmes d'interaction homme-machine, offrant la perspective inédite d'interroger, d'analyser et d'orchestrer ces gisements de données complexes au travers d'un dialogue naturel intuitif.",
        "Cependant, l'adoption effective de l'intelligence artificielle au sein des ERP se heurte à des contraintes de sécurité et de gouvernance strictes. L'utilisation d'API d'IA générative en mode SaaS (hébergées sur des clouds publics tiers) pose des risques majeurs et inacceptables de fuite de données hautement sensibles, d'exposition du secret des affaires et de non-conformité vis-à-vis des cadres légaux relatifs à la protection des données à caractère personnel (tels que la loi n° 2014-038 à Madagascar et le RGPD en Europe). Pour les organisations, la dépendance envers des fournisseurs externes induit également des coûts récurrents indexés sur le volume de jetons (tokens) et une absence totale de maîtrise sur le cycle de vie, la traçabilité et la localisation des données traitées. L'impératif stratégique consiste dès lors à concevoir une solution d'intelligence artificielle « privée » et auto-hébergée (on-premise), opérant en circuit fermé sur l'infrastructure souveraine de l'entreprise.",
        "Au-delà de la souveraineté des modèles, l'intégration d'un agent conversationnel au sein d'un ERP tel qu'Odoo 17 soulève trois défis d'ingénierie majeurs :",
    ]

    intro_bullets = [
        ("La sécurité d'accès et le respect strict du contrôle d'accès basé sur les rôles (RBAC) : ", "L'accès aux enregistrements transactionnels (devis, factures, fiches employés) ne peut souffrir d'aucune élévation implicite de privilèges. Chaque requête de l'IA doit être exécutée strictement sous l'identité de l'utilisateur demandeur (uid) et dans son contexte de cloisonnement multi-sociétés (allowed_company_ids), en respectant à la lettre les listes de contrôle d'accès (ir.model.access) et les règles d'enregistrement dynamiques (ir.rule). Tout contournement par des droits administrateur transversaux (SUPERUSER_ID ou appel à sudo()) est formellement proscrit."),
        ("La prévention des hallucinations et l'exactitude factuelle : ", "Les modèles de langage sont naturellement sujets aux affabulations sémantiques. Dans un contexte de gestion d'entreprise où une erreur de chiffre ou de statut a des répercussions opérationnelles graves, l'agent doit reposer sur des mécanismes rigoureux de récupération de données vérifiables (RAG et requêtes ORM structurées) accompagnées d'une citation formelle des sources."),
        ("L'intégration ergonomique native et la traçabilité dans l'ERP : ", "L'interface conversationnelle doit faire corps avec l'environnement de travail d'Odoo sans perturber les usages quotidiens, tout en assurant une journalisation complète et auditable des interactions pour répondre aux impératifs de conformité."),
    ]

    intro_paragraphs_suite = [
        "Afin de répondre à ces enjeux cruciaux, l'entreprise MIRO K.E. (Silver Partner Odoo) nous a confié la réalisation d'un projet de fin d'études axé sur la conception, la réalisation et l'intégration d'une intelligence artificielle privée et souveraine dans Odoo 17. Dans le cadre de ce mémoire de Master, nous délimitons et validons expérimentalement le cœur opérationnel de ce système, articulé autour de deux composantes logicielles fondamentales :",
    ]

    intro_bullets_axes = [
        ("Le module addon Odoo 17 (« miro_ai ») : ", "conçu nativement avec le framework JavaScript réactif OWL (Odoo Web Library), ce module offre une interface conversationnelle ergonomique intégrée à la barre supérieure (systray) et au chatter d'Odoo. Il intègre un contrôleur sécurisé avec signature cryptographique HMAC, relais du contexte de session utilisateur (uid, company_id, groupes), gestion asynchrone des requêtes par scrutation de tâches (job polling) et journalisation d'audit (miro.ai.log)."),
        ("Le backend d'orchestration et d'intelligence artificielle privée : ", "déployé sous FastAPI et orchestré via LangGraph, il couple un moteur RAG hybride performant (BGE-M3 dense + Rank-BM25 avec fusion Reciprocal Rank Fusion) pour l'interrogation documentaire et un agent de consultation temps réel de l'ORM Odoo sous privilèges stricts sans aucune élévation de privilèges."),
    ]

    intro_paragraphs_persp = [
        "Les dimensions avancées que constituent l'interconnexion universelle par serveur MCP (Model Context Protocol) et l'intelligence de code (Code Intelligence par graphe sémantique) sont quant à elles positionnées et détaillées comme des perspectives structurantes majeures du projet.",
        "Le présent mémoire s'articule autour de trois parties fondamentales :",
    ]

    intro_bullets_structure = [
        ("La Partie I (Présentations) ", "introduit les contextes académique et professionnel en présentant successivement l'École Nationale d'Informatique (ENI) et l'entreprise d'accueil MIRO K.E., avant d'exposer en détail la description du projet, ses objectifs, ses moyens et son chronogramme prévisionnel."),
        ("La Partie II (Analyse et Conception) ", "procède à l'analyse préalable de l'existant au sein de MIRO K.E., puis détaille la conception architecturale globale, la modélisation des cas d'utilisation, les flux de sécurité et les choix méthodologiques et technologiques retenus."),
        ("La Partie III (Réalisation) ", "documente l'environnement d'implémentation, l'intégration du module miro_ai au sein de l'écosystème Odoo 17, ainsi que les protocoles de tests, les benchmarks de performance et l'analyse critique des résultats obtenus."),
    ]

    for p_text in intro_paragraphs:
        p = create_styled_paragraph(doc, p_text)
        partie_1[1].addprevious(p._p)

    for pre, text in intro_bullets:
        p = create_bullet(doc, text, prefix=pre)
        partie_1[1].addprevious(p._p)

    for p_text in intro_paragraphs_suite:
        p = create_styled_paragraph(doc, p_text)
        partie_1[1].addprevious(p._p)

    for pre, text in intro_bullets_axes:
        p = create_bullet(doc, text, prefix=pre)
        partie_1[1].addprevious(p._p)

    for p_text in intro_paragraphs_persp:
        p = create_styled_paragraph(doc, p_text)
        partie_1[1].addprevious(p._p)

    for pre, text in intro_bullets_structure:
        p = create_bullet(doc, text, prefix=pre)
        partie_1[1].addprevious(p._p)

    p_sep = create_styled_paragraph(doc, "")
    partie_1[1].addprevious(p_sep._p)

    print("Introduction Générale successfully inserted.")

    # -------------------------------------------------------------
    # 4. Update CHAPITRE 3: DESCRIPTION DU PROJET
    # -------------------------------------------------------------
    elements_to_delete_chap3 = []
    curr = chap3_h[1].getnext()
    while curr is not None and curr != partie_2[1]:
        elements_to_delete_chap3.append(curr)
        curr = curr.getnext()

    print(f"Removing {len(elements_to_delete_chap3)} legacy elements in Chapter 3...")
    for elem in elements_to_delete_chap3:
        elem.getparent().remove(elem)

    # Insert Chapter 3 elements
    chap3_intro = "Ce chapitre présente de manière exhaustive le cadre fonctionnel, technique et organisationnel du projet. Il expose la formulation précise de la problématique, les objectifs poursuivis, les besoins fonctionnels et non fonctionnels, l'inventaire des ressources humaines, matérielles et logicielles mobilisées, les résultats attendus, les perspectives stratégiques d'extension (notamment le protocole universel MCP et le module de Code Intelligence) ainsi que la planification prévisionnelle des travaux."
    p = create_styled_paragraph(doc, chap3_intro, space_before=6)
    partie_2[1].addprevious(p._p)

    # 3.1 Formulation du projet
    h3_1 = create_subheading(doc, "3.1. Formulation du projet", size_pt=14)
    partie_2[1].addprevious(h3_1._p)

    p3_1_1 = create_styled_paragraph(doc, "Le projet s'intitule : « Conception, réalisation et intégration d’une intelligence artificielle privée dans Odoo 17 ». Il répond à une demande formelle exprimée par la direction technique de la société MIRO K.E., intégrateur et Silver Partner Odoo. L'entreprise constate qu'en dépit de la richesse fonctionnelle d'Odoo 17, les utilisateurs finaux (collaborateurs administratifs, commerciaux, gestionnaires de stocks, techniciens) éprouvent des difficultés régulières à naviguer dans l'arborescence complexe des menus, à synthétiser rapidement des volumes importants de données et à exploiter efficacement la documentation interne.")
    partie_2[1].addprevious(p3_1_1._p)

    p3_1_2 = create_styled_paragraph(doc, "Parallèlement, la direction proscrit formellement le recours aux modèles d'IA hébergés sur le cloud public pour des impératifs stricts de confidentialité, de secret d'affaires et de conformité réglementaire, interdisant tout transfert de données clients, financières ou RH vers des infrastructures tierces.")
    partie_2[1].addprevious(p3_1_2._p)

    p3_1_3 = create_styled_paragraph(doc, "Pour répondre à cette double exigence, le projet adopte une architecture tripartite modulaire :")
    partie_2[1].addprevious(p3_1_3._p)

    arch_tripartite = [
        ("L'addon Odoo 17 « miro_ai » : ", "module d'extension Odoo natif développé avec le framework JavaScript OWL (Odoo Web Library) et Python Odoo. Il s'intègre harmonieusement dans l'interface web (bouton systray, tiroir de chat interactif, rendu markdown et citations de sources), intègre un contrôleur sécurisé avec signature HMAC et clés temporaires, et maintient un journal d'audit complet des interactions (miro.ai.log)."),
        ("Le serveur middleware d'orchestration IA : ", "service backend indépendant développé avec FastAPI et LangGraph, chargé d'analyser les requêtes en langage naturel, de déterminer la route d'exécution optimale et de piloter les agents spécialisés."),
        ("Les moteurs d'exécution spécialisés : ", "d'une part, un moteur RAG hybride (dense BGE-M3 et lexical BM25 avec fusion RRF) assurant la recherche documentaire précise ; d'autre part, un agent de lecture ORM Odoo temps réel interrogeant directement la base de gestion sous les droits et privilèges stricts de l'utilisateur connecté sans recours à sudo()."),
    ]
    for pre, txt in arch_tripartite:
        p_arch = create_bullet(doc, txt, prefix=pre)
        partie_2[1].addprevious(p_arch._p)

    p3_1_4 = create_styled_paragraph(doc, "Le périmètre fonctionnel validé et démontré dans le présent mémoire se focalise sur l'intégration du module miro_ai, l'interrogation documentaire fiable et la consultation temps réel sécurisée de l'ORM Odoo.")
    partie_2[1].addprevious(p3_1_4._p)

    # 3.2 Objectif et besoins de l'utilisateur
    h3_2 = create_subheading(doc, "3.2. Objectif et besoins de l’utilisateur", size_pt=14)
    partie_2[1].addprevious(h3_2._p)

    h3_2_1 = create_subheading(doc, "3.2.1. Objectif général", size_pt=12)
    partie_2[1].addprevious(h3_2_1._p)

    p3_2_1 = create_styled_paragraph(doc, "L'objectif fondamental de ce projet est de concevoir, développer et intégrer dans Odoo 17 une solution d'intelligence artificielle conversationnelle souveraine et privée (miro_ai et son backend dédié), garantissant la stricte confidentialité des données d'entreprise, le respect absolu des privilèges applicatifs RBAC et une assistance performante sur la documentation et les données de gestion.")
    partie_2[1].addprevious(p3_2_1._p)

    h3_2_2 = create_subheading(doc, "3.2.2. Besoins fonctionnels", size_pt=12)
    partie_2[1].addprevious(h3_2_2._p)

    p_bf_intro = create_styled_paragraph(doc, "Les fonctionnalités attendues du système se structurent selon les rôles des utilisateurs :")
    partie_2[1].addprevious(p_bf_intro._p)

    p_bf_user = create_subheading(doc, "Collaborateurs et gestionnaires métier :", size_pt=11)
    partie_2[1].addprevious(p_bf_user._p)

    bf_user_list = [
        "Accéder à l'assistant conversationnel directement depuis n'importe quelle vue de travail d'Odoo 17 (icône systray dans la barre supérieure et volet latéral).",
        "Poser des questions en langage naturel relatives aux procédures internes, politiques d'entreprise et manuels d'utilisation (RAG documentaire).",
        "Consulter l'état des données opérationnelles de l'ERP en temps réel (ex. : devis en attente, factures impayées, niveaux de stock, coordonnées de partenaires).",
        "Visualiser des réponses enrichies avec mise en forme dynamique (Markdown, listes, tableaux), badges de niveau de confiance et citation précise des sources documentaires ou identifiants d'enregistrements.",
        "Consulter l'historique des conversations récentes et reprendre un fil d'échange antérieur.",
    ]
    for b in bf_user_list:
        p_b = create_bullet(doc, b)
        partie_2[1].addprevious(p_b._p)

    p_bf_admin = create_subheading(doc, "Administrateurs du système :", size_pt=11)
    partie_2[1].addprevious(p_bf_admin._p)

    bf_admin_list = [
        "Configurer les paramètres de connexion du module miro_ai dans les paramètres généraux d'Odoo (URL du backend, clé secrète HMAC, délai de réponse timeout).",
        "Activer ou restreindre les domaines de compétences métier de l'assistant (Ventes, Achats, Stocks, Comptabilité, Support technique).",
        "Consulter le journal d'audit complet (miro.ai.log) retraçant les requêtes, les temps de réponse en millisecondes, les sources utilisées et les éventuelles anomalies.",
        "Gérer l'indexation et la mise à jour des bases documentaires vectorielles.",
    ]
    for b in bf_admin_list:
        p_b = create_bullet(doc, b)
        partie_2[1].addprevious(p_b._p)

    h3_2_3 = create_subheading(doc, "3.2.3. Besoins non fonctionnels", size_pt=12)
    partie_2[1].addprevious(h3_2_3._p)

    p_bnf_intro = create_styled_paragraph(doc, "Les exigences qualitatives et techniques indispensables sont les suivantes :")
    partie_2[1].addprevious(p_bnf_intro._p)

    bnf_list = [
        ("Sécurité et gouvernance stricte des données : ", "Principe du moindre privilège : chaque requête ORM doit être exécutée strictement sous l'identité de l'utilisateur connecté (uid), en appliquant les listes de contrôle d'accès (ir.model.access) et les règles d'enregistrement multi-sociétés (ir.rule). Absence totale d'élévation de privilèges : interdiction d'utiliser sudo() ou SUPERUSER_ID pour contourner les droits. Authentification des flux inter-services par HMAC SHA-256."),
        ("Confidentialité et souveraineté locale : ", "Traitement intégral des requêtes et inférence des modèles de langage en local (sur l'infrastructure privée de Miro K.E.). Aucune donnée ne transite par des serveurs tiers."),
        ("Performance et réactivité : ", "Temps de réponse optimisé (< 2 secondes pour l'analyse et la recherche documentaire). Architecture asynchrone par scrutation de tâches (job polling) dans l'addon miro_ai afin de libérer immédiatement les processus web d'Odoo."),
        ("Ergonomie et fidélité à l'écosystème Odoo : ", "Développement de l'interface en JavaScript OWL 2.0 respectant la charte graphique et l'expérience utilisateur standard d'Odoo 17."),
        ("Tolérance aux pannes : ", "Mécanisme de basculement automatique (failover) au niveau de miro_ai en cas d'indisponibilité momentanée d'une instance du backend IA."),
    ]
    for pre, b in bnf_list:
        p_b = create_bullet(doc, b, prefix=pre)
        partie_2[1].addprevious(p_b._p)

    # 3.3 Moyens nécessaires à la réalisation du projet
    h3_3 = create_subheading(doc, "3.3. Moyens nécessaires à la réalisation du projet", size_pt=14)
    partie_2[1].addprevious(h3_3._p)

    p3_3_intro = create_styled_paragraph(doc, "La réalisation de ce travail de recherche et d'ingénierie a mobilisé des ressources humaines, matérielles et logicielles ciblées.")
    partie_2[1].addprevious(p3_3_intro._p)

    # 3.3.1 Moyens humains
    h3_3_1 = create_subheading(doc, "3.3.1. Moyens humains", size_pt=12)
    partie_2[1].addprevious(h3_3_1._p)

    p3_3_1 = create_styled_paragraph(doc, "Le projet implique une collaboration étroite entre l'institution académique et l'entreprise d'accueil, comme l'illustre le Tableau 6.")
    partie_2[1].addprevious(p3_3_1._p)

    # Caption Table 6
    p_cap6 = create_styled_paragraph(doc, "Tableau 6. Moyens humains", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    partie_2[1].addprevious(p_cap6._p)

    headers_tbl6 = ["Rôle / Fonction", "Effectif", "Attribution et responsabilités"]
    data_tbl6 = [
        ["Encadreur pédagogique (ENI)", "1", "Orientation méthodologique, validation scientifique des modèles et suivi académique."],
        ["Encadreur professionnel (Miro K.E.)", "1", "Définition des besoins métier Odoo, cadrage applicatif et accompagnement technique."],
        ["Étudiant-chercheur / Développeur", "1", "Analyse, conception architecturale, développement du module miro_ai, du backend IA et intégration."],
    ]
    tbl6 = create_styled_table(doc, headers_tbl6, data_tbl6, col_widths=[2.2, 0.8, 3.5])
    partie_2[1].addprevious(tbl6._tbl)

    p_space1 = create_styled_paragraph(doc, "")
    partie_2[1].addprevious(p_space1._p)

    # 3.3.2 Moyens matériels
    h3_3_2 = create_subheading(doc, "3.3.2. Moyens matériels", size_pt=12)
    partie_2[1].addprevious(h3_3_2._p)

    p3_3_2 = create_styled_paragraph(doc, "L'exécution locale des modèles de fondation et le développement conjoint de la solution requièrent des équipements informatiques adéquats, détaillés dans le Tableau 7.")
    partie_2[1].addprevious(p3_3_2._p)

    # Caption Table 7
    p_cap7 = create_styled_paragraph(doc, "Tableau 7. Moyens matériels", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    partie_2[1].addprevious(p_cap7._p)

    headers_tbl7 = ["Désignation", "Processeur (CPU)", "Mémoire vive (RAM)", "Stockage", "Accélération GPU", "Affectation"]
    data_tbl7 = [
        ["Station de travail de développement", "Intel Core i7 / AMD Ryzen 7 (8 cœurs, 16 threads)", "32 Go DDR4/DDR5", "SSD NVMe 1 To", "NVIDIA RTX (8 Go VRAM)", "Développement logiciel (miro_ai, FastAPI), tests unitaires, containerisation."],
        ["Serveur d'inférence IA", "Serveur x86_64 multi-cœurs", "64 Go DDR4 ECC", "SSD NVMe 2 To", "NVIDIA RTX / Tesla (16 à 24 Go VRAM)", "Déploiement et exécution des LLM locaux (Ollama/vLLM) et vector store."],
    ]
    tbl7 = create_styled_table(doc, headers_tbl7, data_tbl7, col_widths=[1.3, 1.3, 0.9, 0.9, 1.1, 1.5])
    partie_2[1].addprevious(tbl7._tbl)

    p_space2 = create_styled_paragraph(doc, "")
    partie_2[1].addprevious(p_space2._p)

    # 3.3.3 Moyens logiciels
    h3_3_3 = create_subheading(doc, "3.3.3. Moyens logiciels", size_pt=12)
    partie_2[1].addprevious(h3_3_3._p)

    p3_3_3 = create_styled_paragraph(doc, "L'ensemble de la pile logicielle retenue pour le projet est résumé dans le Tableau 8.")
    partie_2[1].addprevious(p3_3_3._p)

    # Caption Table 8
    p_cap8 = create_styled_paragraph(doc, "Tableau 8. Moyens logiciels", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    partie_2[1].addprevious(p_cap8._p)

    headers_tbl8 = ["N°", "Désignation", "Version", "Rôle et utilisation"]
    data_tbl8 = [
        ["1", "Odoo ERP", "17.0 Community / Enterprise", "Plateforme de gestion intégrée cible hébergeant le module addon miro_ai."],
        ["2", "PostgreSQL", "16.x", "Système de gestion de base de données relationnelle d'Odoo."],
        ["3", "Python", "3.10 / 3.11", "Langage principal du module Odoo miro_ai et du backend d'IA."],
        ["4", "Odoo Web Library (OWL)", "OWL 2.0", "Framework JavaScript réactif utilisé pour l'interface du module miro_ai."],
        ["5", "Ollama / vLLM", "0.3+ / 0.6+", "Moteurs d'inférence et de service pour modèles de fondation locaux."],
        ["6", "Modèles de langage (LLM)", "Qwen2.5 (7B/14B) / Llama 3 (8B)", "Modèles de base pour le raisonnement, l'orchestration et la génération."],
        ["7", "Modèle d'embedding & Rerank", "BGE-M3 (BAAI) / FlagEmbedding", "Génération de représentations vectorielles denses et multi-facettes."],
        ["8", "Framework Agentique", "LangGraph / LangChain", "Modélisation des flux décisionnels sous forme de graphes d'états."],
        ["9", "FastAPI & Pydantic", "0.110+", "Exposition des services d'IA sous forme d'API REST asynchrones."],
        ["10", "Docker & Docker-Compose", "26+", "Conteneurisation et reproductibilité de l'environnement IA."],
        ["11", "Git & GitHub", "2.40+", "Gestion des versions de code source et collaboration."],
        ["12", "Visual Studio Code", "1.94+", "Environnement de développement intégré (IDE) principal."],
    ]
    tbl8 = create_styled_table(doc, headers_tbl8, data_tbl8, col_widths=[0.4, 2.0, 1.4, 3.2])
    partie_2[1].addprevious(tbl8._tbl)

    p_space3 = create_styled_paragraph(doc, "")
    partie_2[1].addprevious(p_space3._p)

    # 3.4 Résultats attendus
    h3_4 = create_subheading(doc, "3.4. Résultats attendus", size_pt=14)
    partie_2[1].addprevious(h3_4._p)

    p3_4_intro = create_styled_paragraph(doc, "Au terme de la mise en œuvre, le projet doit aboutir aux résultats concrets suivants :")
    partie_2[1].addprevious(p3_4_intro._p)

    res_list = [
        ("Module addon Odoo 17 natif (« miro_ai ») : ", "Fourniture d'un module propre, modulaire et directement installable via le gestionnaire d'applications standard d'Odoo, intégrant les vues, la sécurité, les contrôleurs et la traçabilité des échanges."),
        ("Interface utilisateur réactive en OWL : ", "Intégration ergonomique dans le client web d'Odoo via un composant OWL offrant un volet de dialogue accessible en permanence sans altérer l'expérience utilisateur standard."),
        ("RAG hybride haute fidélité : ", "Moteur d'extraction documentaire combinant recherche sémantique dense et lexicale, offrant un taux de rappel supérieur à 90 % et éliminant les hallucinations grâce à la citation rigoureuse des sources documentaires."),
        ("Agent ORM temps réel sécurisé : ", "Restitution instantanée d'indicateurs de gestion tout en garantissant l'étanchéité absolue des données entre utilisateurs et entre sociétés."),
        ("Souveraineté et confidentialité totales : ", "Traitement en circuit fermé attesté par l'absence d'émission de requêtes vers des clouds tiers lors des interactions utilisateur."),
    ]
    for pre, r in res_list:
        p_r = create_bullet(doc, r, prefix=pre)
        partie_2[1].addprevious(p_r._p)

    # 3.5 Perspectives d'évolution du projet (MANDATORY: MCP and Code Intelligence)
    h3_5 = create_subheading(doc, "3.5. Perspectives d'évolution du projet", size_pt=14)
    partie_2[1].addprevious(h3_5._p)

    p3_5_intro = create_styled_paragraph(doc, "Le périmètre opérationnel validé dans le cadre de ce mémoire (l'addon Odoo miro_ai, le RAG documentaire et la lecture ORM sécurisée) pose les fondations d'un écosystème d'IA d'entreprise complet. Deux axes stratégiques majeurs constituent les perspectives d'évolution immédiates du système :")
    partie_2[1].addprevious(p3_5_intro._p)

    # 3.5.1 MCP
    h3_5_1 = create_subheading(doc, "3.5.1. Standardisation des outils via un serveur MCP (Model Context Protocol)", size_pt=12)
    partie_2[1].addprevious(h3_5_1._p)

    p3_5_1_1 = create_styled_paragraph(doc, "L'évolution rapide des architectures agentiques fait émerger des standards d'interconnexion universels. Le protocole MCP (Model Context Protocol), introduit comme standard ouvert par l'industrie, vise à unifier la manière dont les modèles d'IA interagissent avec les sources de données et les outils externes de façon standardisée, modulaire et hautement sécurisée.")
    partie_2[1].addprevious(p3_5_1_1._p)

    p3_5_1_2 = create_styled_paragraph(doc, "Dans cette perspective, il est prévu de développer un serveur MCP natif dédié à Odoo 17. Ce serveur exposera dynamiquement les métadonnées de l'ERP, les schémas de modèles et les outils de consultation sous une interface MCP normalisée. Ainsi, n'importe quel agent ou environnement IA autorisé (qu'il s'agisse d'un client bureautique, d'un IDE assisté par IA ou d'un système multi-agents autonome) pourra interroger et exploiter l'ERP Odoo sans adaptateur propriétaire supplémentaire, tout en conservant le contrôle d'accès strict hérité du contexte de session utilisateur.")
    partie_2[1].addprevious(p3_5_1_2._p)

    # 3.5.2 Code Intelligence
    h3_5_2 = create_subheading(doc, "3.5.2. Module de Code Intelligence et analyse statique des modèles Odoo", size_pt=12)
    partie_2[1].addprevious(h3_5_2._p)

    p3_5_2_1 = create_styled_paragraph(doc, "La seconde perspective majeure réside dans le déploiement du module de Code Intelligence, spécifiquement destiné aux développeurs et administrateurs techniques de Miro K.E. Ce module s'appuie sur une analyse syntaxique abstraite (AST — Abstract Syntax Tree) et sur la construction d'un graphe de connaissances sémantique de plus de 15 000 nœuds et relations, modélisant avec précision l'intégralité de l'écosystème Odoo déployé (modèles Python, champs, vues XML, règles de sécurité ir.model.access, actions de serveur et dépendances croisées de modules).")
    partie_2[1].addprevious(p3_5_2_1._p)

    p3_5_2_2 = create_styled_paragraph(doc, "Grâce à cette représentation structurelle approfondie, le module de Code Intelligence offrira des capacités inédites :")
    partie_2[1].addprevious(p3_5_2_2._p)

    ci_list = [
        ("Explication de workflows complexes : ", "Capacité à restituer et commenter visuellement la chaîne d'héritage d'un modèle et les déclencheurs d'actions automatisées."),
        ("Diagnostic des blocages métier : ", "Identification rapide des conflits d'héritage, des dépendances circulaires ou des restrictions d'accès bloquant un utilisateur."),
        ("Audit statique de sécurité automatisé : ", "Détection préventive d'appels non sécurisés à sudo(), de requêtes SQL brutes non paramétrées ou d'oublis de déclarations de règles d'accès avant la mise en production."),
        ("Assistance à la migration de versions : ", "Accompagnement intelligent lors de la montée de version majeure d'Odoo (par exemple le passage de la v16 vers la v17) en cartographiant les méthodes et champs dépréciés."),
    ]
    for pre, ci in ci_list:
        p_ci = create_bullet(doc, ci, prefix=pre)
        partie_2[1].addprevious(p_ci._p)

    # 3.6 Chronogramme de travail
    h3_6 = create_subheading(doc, "3.6. Chronogramme de travail", size_pt=14)
    partie_2[1].addprevious(h3_6._p)

    p3_6_intro = create_styled_paragraph(doc, "La démarche méthodologique adoptée pour mener à bien le projet s'est échelonnée sur plusieurs phases successives, résumées dans le Tableau 9.")
    partie_2[1].addprevious(p3_6_intro._p)

    # Caption Table 9
    p_cap9 = create_styled_paragraph(doc, "Tableau 9. Chronogramme prévisionnel de réalisation du projet", bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    partie_2[1].addprevious(p_cap9._p)

    headers_tbl9 = ["Phase", "Intitulé des travaux", "Période indicative", "Livrables associés"]
    data_tbl9 = [
        ["Phase 1", "Analyse des besoins, état de l'art et étude de faisabilité de l'IA locale", "Mois 1", "Rapport d'analyse préalable, benchmark comparatif des LLM et architectures RAG."],
        ["Phase 2", "Conception architecturale, modélisation de la sécurité et des flux", "Mois 2", "Spécifications techniques, diagrammes UML, matrice des flux et schéma de sécurité."],
        ["Phase 3", "Développement du moteur RAG hybride et de l'agent ORM sécurisé", "Mois 3", "Pipeline d'indexation/rerank BGE-M3, endpoints FastAPI, connecteurs ORM sécurisés."],
        ["Phase 4", "Développement du module Odoo 17 (miro_ai) et intégration UI OWL", "Mois 4", "Addon Odoo miro_ai (composants OWL, contrôleurs, journal d'audit miro.ai.log)."],
        ["Phase 5", "Campagne de tests, benchmarking des performances et évaluation", "Mois 5", "Jeux d'essais, métriques de rappel/latence, protocole d'évaluation et validation fonctionnelle."],
        ["Phase 6", "Rédaction finale du mémoire de Master et préparation de la soutenance", "Mois 6", "Manuscrit du mémoire de Master ENI complet, supports de présentation."],
    ]
    tbl9 = create_styled_table(doc, headers_tbl9, data_tbl9, col_widths=[0.8, 2.3, 1.2, 2.7])
    partie_2[1].addprevious(tbl9._tbl)

    p_space4 = create_styled_paragraph(doc, "")
    partie_2[1].addprevious(p_space4._p)

    print("Chapter 3 successfully inserted.")

    # Save document
    print(f"Saving document to: {doc_path}...")
    doc.save(doc_path)
    print("Document successfully saved!")

if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else '/home/user/Downloads/MASTER_ENI.docx'
    update_document(target)
